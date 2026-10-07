#!/usr/bin/env python3
# -*- coding: utf-8 -*-
r"""
TRACE Phase 0B
==============
Minimal incremental acquisition + reference freeze + workbook/schema audit.

Project root:
    D:\RNA\Trace

This script:
1) downloads only small processed/source-data assets needed by the frozen design;
2) never downloads FASTQ/BAM/raw sequencing;
3) freezes the existing MANE GTF only if duplicate local copies are byte-identical;
4) audits every XLSX workbook under 01_source without modifying it;
5) extracts the 2026 DRS supplementary ZIP if available;
6) writes machine-readable reports for the next harmonization step.

It never overwrites a conflicting existing file.
"""

from __future__ import annotations

import csv
import hashlib
import html
import json
import re
import shutil
import time
import traceback
import urllib.parse
import urllib.request
import zipfile
from pathlib import Path
from typing import Dict, List, Tuple

ROOT = Path(r"D:\RNA\Trace")
META = ROOT / "00_meta"
SRC = ROOT / "01_source"
REF = ROOT / "02_reference"
LOGS = ROOT / "logs"

for p in (META, SRC, REF, LOGS):
    p.mkdir(parents=True, exist_ok=True)

DOWNLOADS = [
    {
        "asset_id": "BID_A549_WT",
        "purpose": "Aim1 independent A549 replication",
        "url": "https://ftp.ncbi.nlm.nih.gov/geo/series/GSE179nnn/GSE179798/suppl/GSE179798_A549_mRNA_WT_BID-seq.xlsx",
        "dest": SRC / "BID" / "GSE179798_A549_mRNA_WT_BID-seq.xlsx",
        "kind": "xlsx",
    },
    {
        "asset_id": "BID_supplementary_tables",
        "purpose": "BID NN-Psi-NN calibration parameters; Supplementary Tables 1-23",
        "url": "https://media.springernature.com/original/springer-static/esm/art%3A10.1038%2Fs41587-022-01505-w/MediaObjects/41587_2022_1505_MOESM3_ESM.xlsx",
        "dest": SRC / "BID" / "BID_supplementary_tables.xlsx",
        "kind": "xlsx",
    },
    {
        "asset_id": "BID_source_data_fig1",
        "purpose": "BID Fig.1 synthetic-context source data",
        "url": "https://media.springernature.com/original/springer-static/esm/art%3A10.1038%2Fs41587-022-01505-w/MediaObjects/41587_2022_1505_MOESM14_ESM.xlsx",
        "dest": SRC / "BID" / "BID_source_data_fig1.xlsx",
        "kind": "xlsx",
    },
    {
        "asset_id": "BACS_source_data_fig1",
        "purpose": "BACS Fig.1 synthetic-context source data",
        "url": "https://media.springernature.com/original/springer-static/esm/art%3A10.1038%2Fs41592-024-02439-8/MediaObjects/41592_2024_2439_MOESM4_ESM.xlsx",
        "dest": SRC / "BACS" / "BACS_source_data_fig1.xlsx",
        "kind": "xlsx",
    },
]

DRS_ARTICLE = "https://pmc.ncbi.nlm.nih.gov/articles/PMC13096809/"
DRS_DIR = SRC / "DRS_2026"
DRS_ZIP = DRS_DIR / "gkag353_supplemental_files.zip"
DRS_EXTRACTED = DRS_DIR / "extracted"

USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124 Safari/537.36"
)

download_rows: List[Dict] = []
log_lines: List[str] = []


def log(msg: str) -> None:
    stamp = time.strftime("%Y-%m-%d %H:%M:%S")
    line = f"[{stamp}] {msg}"
    print(line)
    log_lines.append(line)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def validate_zip_container(path: Path) -> Tuple[bool, str]:
    try:
        with zipfile.ZipFile(path, "r") as zf:
            bad = zf.testzip()
            if bad:
                return False, f"corrupt member: {bad}"
        return True, "ok"
    except Exception as e:
        return False, repr(e)


def safe_download(asset_id: str, purpose: str, url: str, dest: Path, kind: str) -> None:
    dest.parent.mkdir(parents=True, exist_ok=True)

    if dest.exists():
        valid, why = (validate_zip_container(dest) if kind in {"xlsx", "zip"} else (True, "ok"))
        download_rows.append({
            "asset_id": asset_id,
            "purpose": purpose,
            "status": "PRESENT_EXISTING",
            "url": url,
            "local_path": str(dest),
            "size_bytes": dest.stat().st_size,
            "sha256": sha256(dest),
            "validation": "PASS" if valid else "FAIL",
            "notes": why,
        })
        log(f"{asset_id}: already present; not overwritten.")
        return

    tmp = dest.with_suffix(dest.suffix + ".part")
    if tmp.exists():
        tmp.unlink()

    req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    try:
        log(f"Downloading {asset_id} ...")
        with urllib.request.urlopen(req, timeout=180) as r, tmp.open("wb") as f:
            shutil.copyfileobj(r, f)
        tmp.replace(dest)

        valid, why = (validate_zip_container(dest) if kind in {"xlsx", "zip"} else (True, "ok"))
        download_rows.append({
            "asset_id": asset_id,
            "purpose": purpose,
            "status": "DOWNLOADED",
            "url": url,
            "local_path": str(dest),
            "size_bytes": dest.stat().st_size,
            "sha256": sha256(dest),
            "validation": "PASS" if valid else "FAIL",
            "notes": why,
        })
    except Exception as e:
        if tmp.exists():
            try:
                tmp.unlink()
            except Exception:
                pass
        download_rows.append({
            "asset_id": asset_id,
            "purpose": purpose,
            "status": "DOWNLOAD_FAILED",
            "url": url,
            "local_path": str(dest),
            "size_bytes": "",
            "sha256": "",
            "validation": "FAIL",
            "notes": repr(e),
        })
        log(f"{asset_id}: download failed: {e}")


def resolve_and_download_drs() -> None:
    DRS_DIR.mkdir(parents=True, exist_ok=True)
    if DRS_ZIP.exists():
        valid, why = validate_zip_container(DRS_ZIP)
        download_rows.append({
            "asset_id": "DRS_2026_supplement",
            "purpose": "DRS sites + protein/TPM/TE + housekeeping controls",
            "status": "PRESENT_EXISTING",
            "url": DRS_ARTICLE,
            "local_path": str(DRS_ZIP),
            "size_bytes": DRS_ZIP.stat().st_size,
            "sha256": sha256(DRS_ZIP),
            "validation": "PASS" if valid else "FAIL",
            "notes": why,
        })
        if valid:
            extract_drs()
        return

    candidates = []

    try:
        log("Resolving 2026 DRS supplementary ZIP from PMC ...")
        req = urllib.request.Request(DRS_ARTICLE, headers={"User-Agent": USER_AGENT})
        with urllib.request.urlopen(req, timeout=120) as r:
            page = r.read().decode("utf-8", errors="replace")

        patterns = [
            r'href=["\']([^"\']*gkag353_supplemental_files\.zip[^"\']*)["\']',
            r'href=["\']([^"\']*gkag353[^"\']*supplement[^"\']*\.zip[^"\']*)["\']',
        ]
        for pat in patterns:
            for m in re.finditer(pat, page, flags=re.I):
                href = html.unescape(m.group(1))
                url = urllib.parse.urljoin(DRS_ARTICLE, href)
                if url not in candidates:
                    candidates.append(url)
    except Exception as e:
        log(f"PMC HTML resolution warning: {e}")

    candidates.extend([
        "https://pmc.ncbi.nlm.nih.gov/articles/PMC13096809/bin/gkag353_supplemental_files.zip",
        "https://pmc.ncbi.nlm.nih.gov/articles/instance/13096809/bin/gkag353_supplemental_files.zip",
        "https://www.ncbi.nlm.nih.gov/pmc/articles/PMC13096809/bin/gkag353_supplemental_files.zip",
    ])

    seen = set()
    last_error = ""
    for url in candidates:
        if url in seen:
            continue
        seen.add(url)
        tmp = DRS_ZIP.with_suffix(".zip.part")
        if tmp.exists():
            try:
                tmp.unlink()
            except Exception:
                pass
        try:
            log("Trying DRS supplement attachment ...")
            req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
            with urllib.request.urlopen(req, timeout=180) as r, tmp.open("wb") as f:
                shutil.copyfileobj(r, f)
            tmp.replace(DRS_ZIP)
            valid, why = validate_zip_container(DRS_ZIP)
            if not valid:
                last_error = why
                try:
                    DRS_ZIP.unlink()
                except Exception:
                    pass
                continue

            download_rows.append({
                "asset_id": "DRS_2026_supplement",
                "purpose": "DRS sites + protein/TPM/TE + housekeeping controls",
                "status": "DOWNLOADED",
                "url": url,
                "local_path": str(DRS_ZIP),
                "size_bytes": DRS_ZIP.stat().st_size,
                "sha256": sha256(DRS_ZIP),
                "validation": "PASS",
                "notes": "ok",
            })
            extract_drs()
            return
        except Exception as e:
            last_error = repr(e)
            if tmp.exists():
                try:
                    tmp.unlink()
                except Exception:
                    pass

    download_rows.append({
        "asset_id": "DRS_2026_supplement",
        "purpose": "DRS sites + protein/TPM/TE + housekeeping controls",
        "status": "MANUAL_DOWNLOAD_REQUIRED",
        "url": DRS_ARTICLE,
        "local_path": str(DRS_ZIP),
        "size_bytes": "",
        "sha256": "",
        "validation": "FAIL",
        "notes": last_error or "Could not resolve/download attachment.",
    })
    log("DRS supplement could not be downloaded automatically; other Phase 0B work will continue.")


def extract_drs() -> None:
    DRS_EXTRACTED.mkdir(parents=True, exist_ok=True)
    marker = DRS_EXTRACTED / ".trace_extracted_ok"
    current_hash = sha256(DRS_ZIP)

    if marker.exists() and marker.read_text(encoding="utf-8", errors="ignore").strip() == current_hash:
        log("DRS supplementary ZIP already extracted.")
        return

    with zipfile.ZipFile(DRS_ZIP, "r") as zf:
        zf.extractall(DRS_EXTRACTED)
    marker.write_text(current_hash, encoding="utf-8")
    log(f"DRS supplement extracted to {DRS_EXTRACTED}")


def write_tsv(path: Path, rows: List[Dict], fields: List[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, delimiter="\t", extrasaction="ignore")
        w.writeheader()
        for row in rows:
            w.writerow(row)


def freeze_mane_reference() -> Dict:
    result = {
        "status": "NO_REFERENCE_CANDIDATES",
        "selected_source": "",
        "trace_path": "",
        "sha256": "",
        "candidate_count": 0,
        "all_candidates_identical": "",
        "notes": "",
    }

    candidates_file = META / "reference_candidates.tsv"
    if not candidates_file.exists():
        result["notes"] = "reference_candidates.tsv is missing."
        return result

    candidates = []
    with candidates_file.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f, delimiter="\t"):
            p = Path(row.get("path", ""))
            if p.exists() and "mane" in p.name.lower() and p.name.lower().endswith((".gtf", ".gtf.gz")):
                candidates.append(p)

    result["candidate_count"] = len(candidates)
    if not candidates:
        result["notes"] = "No local MANE GTF found."
        return result

    hashes = [(p, sha256(p)) for p in candidates]
    unique_hashes = sorted(set(h for _, h in hashes))
    result["all_candidates_identical"] = (len(unique_hashes) == 1)

    if len(unique_hashes) != 1:
        result["status"] = "REFERENCE_CONFLICT"
        result["notes"] = "; ".join(f"{p}={h}" for p, h in hashes)
        return result

    preferred = next(
        (p for p, _ in hashes if "\\phase1\\" in str(p).lower()),
        hashes[0][0]
    )
    dst = REF / preferred.name

    if dst.exists():
        if sha256(dst) != unique_hashes[0]:
            result["status"] = "TRACE_REFERENCE_CONFLICT"
            result["selected_source"] = str(preferred)
            result["trace_path"] = str(dst)
            result["sha256"] = sha256(dst)
            result["notes"] = "Existing Trace reference differs; not overwritten."
            return result
        status = "PRESENT_VERIFIED"
    else:
        shutil.copy2(preferred, dst)
        status = "FROZEN_VERIFIED"

    result.update({
        "status": status,
        "selected_source": str(preferred),
        "trace_path": str(dst),
        "sha256": unique_hashes[0],
        "notes": "Duplicate local MANE GTF copies were byte-identical.",
    })
    return result


def audit_workbooks():
    try:
        import openpyxl
    except Exception as e:
        return [], [], [{
            "file": "",
            "sheet": "",
            "hit_type": "DEPENDENCY_MISSING",
            "detail": f"openpyxl import failed: {e}",
        }]

    schema_rows = []
    preview_blocks = []
    target_hits = []

    xlsx_files = sorted([p for p in SRC.rglob("*.xlsx") if p.is_file()])

    target_terms = {
        "calibration": ["NNΨNN", "NNUNN", "deletion", "conversion", "dropout", "motif"],
        "drs_sites": ["chrom", "position", "mismatch", "occupancy", "psi"],
        "functional": ["protein", "TPM", "translation efficiency", "TE"],
        "cell_lines": ["HeLa", "A549"],
        "writer": ["TRUB1", "PUS7", "PUS1", "DKC1"],
    }

    for path in xlsx_files:
        rel = path.relative_to(ROOT)
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except Exception as e:
            schema_rows.append({
                "file": str(rel), "sheet": "", "rows": "", "columns": "",
                "header_guess_row": "", "header_guess": "",
                "status": "OPEN_FAILED", "notes": repr(e),
            })
            continue

        preview_blocks.append("=" * 100)
        preview_blocks.append(f"FILE: {rel}")
        preview_blocks.append(f"SHEETS: {wb.sheetnames}")

        for ws in wb.worksheets:
            max_r = min(ws.max_row or 1, 15)
            max_c = min(ws.max_column or 1, 20)
            sample = [
                list(row)
                for row in ws.iter_rows(
                    min_row=1, max_row=max_r, min_col=1, max_col=max_c, values_only=True
                )
            ]

            best_idx = 0
            best_score = -1
            for i, row in enumerate(sample):
                nonempty = sum(v is not None and str(v).strip() != "" for v in row)
                string_bonus = sum(isinstance(v, str) and v.strip() != "" for v in row)
                total = nonempty + 0.25 * string_bonus
                if total > best_score:
                    best_score = total
                    best_idx = i

            header = sample[best_idx] if sample else []
            header_clean = [str(v).strip() if v is not None else "" for v in header]
            header_text = " | ".join(x for x in header_clean if x)

            schema_rows.append({
                "file": str(rel),
                "sheet": ws.title,
                "rows": ws.max_row,
                "columns": ws.max_column,
                "header_guess_row": best_idx + 1 if sample else "",
                "header_guess": header_text[:1200],
                "status": "OK",
                "notes": "",
            })

            preview_blocks.append("-" * 100)
            preview_blocks.append(f"SHEET: {ws.title}  rows={ws.max_row} cols={ws.max_column}")
            for i, row in enumerate(sample[:8], start=1):
                vals = ["" if v is None else str(v) for v in row[:12]]
                preview_blocks.append(f"{i:>3}: " + " || ".join(vals))

            haystack = (
                ws.title + " " + header_text + " " +
                " ".join(
                    " ".join("" if v is None else str(v) for v in row[:20])
                    for row in sample[:10]
                )
            ).lower()

            for hit_type, terms in target_terms.items():
                hits = [term for term in terms if term.lower() in haystack]
                if hits:
                    target_hits.append({
                        "file": str(rel),
                        "sheet": ws.title,
                        "hit_type": hit_type,
                        "detail": ", ".join(hits),
                    })

        wb.close()

    return schema_rows, preview_blocks, target_hits


def main():
    print("=" * 88)
    print("TRACE PHASE 0B — INCREMENTAL DATA + SCHEMA AUDIT")
    print("=" * 88)

    for item in DOWNLOADS:
        safe_download(
            item["asset_id"], item["purpose"], item["url"], item["dest"], item["kind"]
        )

    resolve_and_download_drs()

    log("Freezing local MANE annotation reference ...")
    reference_result = freeze_mane_reference()
    (META / "reference_freeze.json").write_text(
        json.dumps(reference_result, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    log("Auditing XLSX workbook schemas ...")
    schema_rows, preview_blocks, target_hits = audit_workbooks()

    download_fields = [
        "asset_id", "purpose", "status", "url", "local_path",
        "size_bytes", "sha256", "validation", "notes",
    ]
    write_tsv(META / "phase0b_download_manifest.tsv", download_rows, download_fields)

    schema_fields = [
        "file", "sheet", "rows", "columns", "header_guess_row",
        "header_guess", "status", "notes",
    ]
    write_tsv(META / "workbook_schema.tsv", schema_rows, schema_fields)

    hit_fields = ["file", "sheet", "hit_type", "detail"]
    write_tsv(META / "target_table_hits.tsv", target_hits, hit_fields)

    (META / "workbook_preview.txt").write_text(
        "\n".join(preview_blocks), encoding="utf-8"
    )

    drs_rows = []
    if DRS_EXTRACTED.exists():
        for p in sorted(DRS_EXTRACTED.rglob("*")):
            if not p.is_file() or p.name == ".trace_extracted_ok":
                continue
            drs_rows.append({
                "relative_path": str(p.relative_to(DRS_EXTRACTED)),
                "size_bytes": p.stat().st_size,
                "suffix": p.suffix.lower(),
            })
    write_tsv(
        META / "drs_extracted_inventory.tsv",
        drs_rows,
        ["relative_path", "size_bytes", "suffix"]
    )

    failed = [r for r in download_rows if r["validation"] != "PASS"]
    summary = {
        "phase": "0B",
        "status": "PHASE0B_AUDIT_COMPLETE",
        "downloads_total": len(download_rows),
        "downloads_pass": len(download_rows) - len(failed),
        "downloads_attention": [r["asset_id"] for r in failed],
        "reference_freeze": reference_result,
        "xlsx_files_audited": len(set(r["file"] for r in schema_rows if r["file"])),
        "xlsx_sheets_audited": len([r for r in schema_rows if r["sheet"]]),
        "target_table_hits": len(target_hits),
        "drs_extracted_files": len(drs_rows),
        "genome_fasta_status": "DEFERRED_UNTIL_SCHEMA_REVIEW",
        "scientific_gate": (
            "Do not construct canonical loci yet. First review workbook_schema, "
            "target_table_hits, and DRS extracted contents; then freeze field mappings."
        ),
    }
    (META / "phase0b_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    (LOGS / "phase0b_console.log").write_text(
        "\n".join(log_lines) + "\n", encoding="utf-8"
    )

    print()
    print("=" * 88)
    print("PHASE 0B AUDIT COMPLETE")
    print("=" * 88)
    print(f"Downloads/assets checked: {summary['downloads_total']}")
    print(f"Passed:                   {summary['downloads_pass']}")
    print(f"Attention:                {summary['downloads_attention']}")
    print(f"Reference:                {reference_result['status']}")
    print(f"XLSX files audited:       {summary['xlsx_files_audited']}")
    print(f"Sheets audited:           {summary['xlsx_sheets_audited']}")
    print(f"Target-table hits:        {summary['target_table_hits']}")
    print(f"DRS extracted files:      {summary['drs_extracted_files']}")
    print()
    print(r"Return these files from D:\RNA\Trace\00_meta:")
    print("  phase0b_download_manifest.tsv")
    print("  reference_freeze.json")
    print("  workbook_schema.tsv")
    print("  target_table_hits.tsv")
    print("  workbook_preview.txt")
    print("  drs_extracted_inventory.tsv")
    print("  phase0b_summary.json")
    print()
    print("Genome FASTA is intentionally NOT downloaded yet.")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception:
        err = traceback.format_exc()
        try:
            LOGS.mkdir(parents=True, exist_ok=True)
            (LOGS / "phase0b_FATAL_ERROR.txt").write_text(err, encoding="utf-8")
        except Exception:
            pass
        print("\nFATAL ERROR\n")
        print(err)
        print(r"Send D:\RNA\Trace\logs\phase0b_FATAL_ERROR.txt back to ChatGPT.")
        try:
            input("\nPress Enter to close...")
        except Exception:
            pass
        raise
