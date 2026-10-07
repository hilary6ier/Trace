#!/usr/bin/env python3
# -*- coding: utf-8 -*-
r"""
TRACE Phase 0C — formal source normalization
=============================================

Scientific scope
----------------
This step creates provenance-preserving NORMALIZED SOURCE TABLES only.
It does NOT:
- create assay-exclusive labels,
- treat missing calls as negatives,
- infer "true Psi",
- perform arbitrary +/-k matching,
- train any model.

It explicitly implements the frozen measurement-science design:
Calibration -> Measurement -> Evidence -> Inference

Project root:
    D:\RNA\Trace

Requirements:
    Python 3
    openpyxl >= 3.1

Outputs:
    D:\RNA\Trace\03_harmonized\normalized_sources\
    D:\RNA\Trace\03_harmonized\synthetic_context_calibration.tsv
    D:\RNA\Trace\00_meta\phase0c_*.{tsv,json,txt}
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import shutil
import statistics
import traceback
import zipfile
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

try:
    import openpyxl
except Exception as e:
    raise SystemExit(
        "openpyxl is required. In the same Anaconda environment run:\n"
        "  conda install -y openpyxl\n"
        f"Original import error: {e}"
    )

ROOT = Path(r"D:\RNA\Trace")
META = ROOT / "00_meta"
SOURCE = ROOT / "01_source"
REF = ROOT / "02_reference"
HARM = ROOT / "03_harmonized"
NORM = HARM / "normalized_sources"
QUAR = HARM / "quarantine"
LOGS = ROOT / "logs"

for p in [META, HARM, NORM, QUAR, LOGS]:
    p.mkdir(parents=True, exist_ok=True)

# ---------------------------------------------------------------------
# Frozen scientific contracts from validated Stage0 + primary papers
# ---------------------------------------------------------------------

CHEMISTRY = {
    "BID": "bisulfite_deletion",
    "PRAISE": "bisulfite_deletion",
    "BACS": "cyclization_conversion",
    "ELAP": "enzymatic_labeling_RT_stop",
    "DRS": "native_direct_RNA",
}

EXPECTED_COUNTS = {
    "BID_HEK293T": 543,
    "BID_HeLa": 575,
    "BID_A549": 922,
    "ELAP_HEK293T": 5731,
    "ELAP_HeLa": 5367,
    "BACS_HeLa": 1335,
    "PRAISE_HEK293T": 2209,
}

COORDINATE_CONTRACT = {
    "BID": {
        "assembly": "GRCh38/hg38",
        "source_semantics": "1-based genomic nucleotide",
        "normalized_pos1": "source pos",
        "status": "validated_in_Stage0",
    },
    "BACS": {
        "assembly": "GRCh38/hg38",
        "source_semantics": "1-based genomic nucleotide",
        "normalized_pos1": "source pos",
        "status": "validated_in_Stage0",
    },
    "ELAP": {
        "assembly": "GRCh38/hg38",
        "source_semantics": "one-base BED-like [start,end)",
        "normalized_pos1": "end",
        "status": "validated_in_Stage0",
    },
    "PRAISE": {
        "assembly": "GRCh38/hg38",
        "source_semantics": "published chr_site may be a nucleotide or interval",
        "normalized_pos1": "NOT_FORCED",
        "status": "interval/equivalence handling required",
    },
    "DRS": {
        "assembly": "to_be_read_from_published_table",
        "source_semantics": "preserve published coordinate fields verbatim in Phase0C",
        "normalized_pos1": "NOT_FORCED",
        "status": "held-out platform; exact mapping deferred until table semantics validated",
    },
}

# Known source files.
PATHS = {
    "BID_HEK293T": SOURCE / "BID" / "GSE179798_HEK293T_mRNA_WT_BID-seq.xlsx",
    "BID_HeLa": SOURCE / "BID" / "GSE179798_HeLa_mRNA_WT_BID-seq.xlsx",
    "BID_A549": SOURCE / "BID" / "GSE179798_A549_mRNA_WT_BID-seq.xlsx",
    "BID_shControl": SOURCE / "BID" / "GSE179798_HeLa_mRNA_shControl_BID-seq.xlsx",
    "BID_supp": SOURCE / "BID" / "BID_supplementary_tables.xlsx",
    "BID_fig1": SOURCE / "BID" / "BID_source_data_fig1.xlsx",
    "ELAP_HEK293T": SOURCE / "ELAP" / "GSE236530_ELAP-HEK-all-update-1.xlsx",
    "ELAP_HeLa": SOURCE / "ELAP" / "GSE236530_ELAP-HeLa-all-update-1.xlsx",
    "ELAP_DKC1": SOURCE / "ELAP" / "GSE236530_DKC1-knockdown-enrichment-change.xlsx",
    "ELAP_supp": SOURCE / "ELAP" / "ELAP_supplementary_data.xlsx",
    "ELAP_source": SOURCE / "ELAP" / "ELAP_source_data.xlsx",
    "BACS_HeLa": SOURCE / "BACS" / "GSE241849_Supplementary_Table_3_10.xlsx",
    "BACS_fig1": SOURCE / "BACS" / "BACS_source_data_fig1.xlsx",
    "PRAISE": SOURCE / "PRAISE" / "41589_2023_1304_MOESM3_ESM.xlsx",
    "DRS_zip": SOURCE / "DRS_2026" / "gkag353_supplemental_files.zip",
    "MANE": REF / "MANE.GRCh38.v1.5.refseq_genomic.gtf.gz",
}

DRS_EXTRACT = SOURCE / "DRS_2026" / "extracted"

qc_rows: List[Dict[str, Any]] = []
error_rows: List[Dict[str, Any]] = []
schema_rows: List[Dict[str, Any]] = []
manifest_rows: List[Dict[str, Any]] = []
calibration_mapping_rows: List[Dict[str, Any]] = []
column_dictionary_rows: List[Dict[str, Any]] = []


# ---------------------------------------------------------------------
# General utilities
# ---------------------------------------------------------------------

def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def safe_str(x: Any) -> str:
    if x is None:
        return ""
    if isinstance(x, float) and math.isnan(x):
        return ""
    return str(x).strip()


def to_float(x: Any) -> Optional[float]:
    if x is None or x == "":
        return None
    if isinstance(x, (int, float)):
        try:
            y = float(x)
            return None if math.isnan(y) else y
        except Exception:
            return None
    s = str(x).strip().replace(",", "")
    s = s.replace("%", "")
    try:
        return float(s)
    except Exception:
        return None


def clean_col(name: Any) -> str:
    s = safe_str(name)
    if not s:
        return "unnamed"
    s = s.replace("Ψ", "psi")
    s = s.replace("′", "").replace("’", "")
    s = re.sub(r"\s+", "_", s)
    s = re.sub(r"[^A-Za-z0-9_.()+%-]+", "_", s)
    s = re.sub(r"_+", "_", s).strip("_")
    return s or "unnamed"


def make_unique_columns(cols: Sequence[Any]) -> Tuple[List[str], List[Tuple[str, str]]]:
    out = []
    seen: Dict[str, int] = {}
    mapping = []
    for raw in cols:
        base = clean_col(raw)
        n = seen.get(base, 0)
        seen[base] = n + 1
        c = base if n == 0 else f"{base}__dup{n+1}"
        out.append(c)
        mapping.append((c, safe_str(raw)))
    return out, mapping


def write_tsv(path: Path, rows: Sequence[Dict[str, Any]], fields: Optional[Sequence[str]] = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    rows = list(rows)
    if fields is None:
        ordered: List[str] = []
        seen = set()
        for r in rows:
            for k in r.keys():
                if k not in seen:
                    seen.add(k)
                    ordered.append(k)
        fields = ordered
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(fields), delimiter="\t", extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow({k: "" if r.get(k) is None else r.get(k) for k in fields})


def add_qc(check_id: str, dataset: str, status: str, observed: Any = "", expected: Any = "", notes: str = "") -> None:
    qc_rows.append({
        "check_id": check_id,
        "dataset": dataset,
        "status": status,
        "observed": observed,
        "expected": expected,
        "notes": notes,
    })


def add_error(stage: str, dataset: str, message: str) -> None:
    error_rows.append({"stage": stage, "dataset": dataset, "message": message})


def list_sheets(path: Path) -> List[str]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    names = list(wb.sheetnames)
    wb.close()
    return names


def read_matrix(path: Path, sheet: str, max_rows: Optional[int] = None) -> List[List[Any]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    matrix = []
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        matrix.append(list(row))
        if max_rows is not None and i >= max_rows:
            break
    wb.close()
    return matrix


def detect_header_row(matrix: Sequence[Sequence[Any]], keywords: Sequence[str], max_scan: int = 25) -> int:
    kws = [k.lower() for k in keywords]
    best = (float("-inf"), 0)
    for i, row in enumerate(matrix[:max_scan]):
        vals = [safe_str(x) for x in row]
        text = " ".join(vals).lower()
        nonempty = sum(bool(v) for v in vals)
        strings = sum(bool(v) and not re.fullmatch(r"[-+]?\d+(\.\d+)?", v) for v in vals)
        hits = sum(k in text for k in kws)
        score = 5 * hits + nonempty + 0.25 * strings
        if score > best[0]:
            best = (score, i)
    return best[1]


def records_from_matrix(matrix: Sequence[Sequence[Any]], header_idx: int) -> Tuple[List[Dict[str, Any]], List[str], List[Tuple[str, str]]]:
    raw_header = list(matrix[header_idx])
    cols, mapping = make_unique_columns(raw_header)
    records = []
    for ridx, row in enumerate(matrix[header_idx + 1 :], start=header_idx + 2):
        vals = list(row) + [None] * max(0, len(cols) - len(row))
        vals = vals[:len(cols)]
        if not any(safe_str(v) for v in vals):
            continue
        rec = {c: v for c, v in zip(cols, vals)}
        rec["_excel_row"] = ridx
        records.append(rec)
    return records, cols, mapping


def read_records(path: Path, sheet: str, header_idx: Optional[int], keywords: Sequence[str]) -> Tuple[List[Dict[str, Any]], List[str], int]:
    matrix = read_matrix(path, sheet)
    if header_idx is None:
        header_idx = detect_header_row(matrix, keywords)
    recs, cols, mapping = records_from_matrix(matrix, header_idx)
    for cleaned, original in mapping:
        column_dictionary_rows.append({
            "source_file": path.name,
            "source_sheet": sheet,
            "clean_column": cleaned,
            "original_column": original,
        })
    schema_rows.append({
        "source_file": path.name,
        "source_sheet": sheet,
        "header_row_1based": header_idx + 1,
        "data_rows": len(recs),
        "n_columns": len(cols),
        "columns": " | ".join(cols),
    })
    return recs, cols, header_idx


def first_existing(rec: Dict[str, Any], candidates: Sequence[str]) -> Any:
    for c in candidates:
        if c in rec and safe_str(rec[c]) != "":
            return rec[c]
    return None


def normalize_chrom(x: Any) -> str:
    s = safe_str(x)
    if not s:
        return ""
    if s.lower().startswith("chr"):
        return "chr" + s[3:]
    if re.fullmatch(r"(?:[0-9]+|X|Y|M|MT)", s, flags=re.I):
        if s.upper() == "MT":
            s = "M"
        return "chr" + s
    return s


def normalize_strand(x: Any) -> str:
    s = safe_str(x)
    if s in {"+", "-"}:
        return s
    return s


def motif5(x: Any) -> str:
    s = safe_str(x).upper()
    s = s.replace("Ψ", "U").replace("T", "U")
    s = re.sub(r"[^ACGU]", "", s)
    return s if len(s) == 5 else ""


def local_urun_len_from_5mer(m: str) -> Optional[int]:
    m = motif5(m)
    if not m or m[2] != "U":
        return None
    l = 2
    r = 2
    while l - 1 >= 0 and m[l - 1] == "U":
        l -= 1
    while r + 1 < len(m) and m[r + 1] == "U":
        r += 1
    return r - l + 1


def percentile_rank(values: Sequence[Optional[float]]) -> List[Optional[float]]:
    idx_vals = [(i, v) for i, v in enumerate(values) if v is not None]
    n = len(idx_vals)
    out: List[Optional[float]] = [None] * len(values)
    if n == 0:
        return out
    sorted_vals = sorted(v for _, v in idx_vals)
    # average rank for ties
    pos = defaultdict(list)
    for j, v in enumerate(sorted_vals, start=1):
        pos[v].append(j)
    rank = {v: statistics.mean(js) for v, js in pos.items()}
    denom = max(n - 1, 1)
    for i, v in idx_vals:
        out[i] = (rank[v] - 1) / denom
    return out


def extract_raw_columns(rec: Dict[str, Any], exclude: Sequence[str] = ("_excel_row",)) -> Dict[str, Any]:
    return {f"raw__{k}": v for k, v in rec.items() if k not in exclude}


# ---------------------------------------------------------------------
# Asset preflight
# ---------------------------------------------------------------------

def preflight() -> None:
    required = [
        "BID_HEK293T", "BID_HeLa", "BID_A549",
        "ELAP_HEK293T", "ELAP_HeLa",
        "BACS_HeLa", "PRAISE",
        "BID_supp", "BACS_fig1", "ELAP_source",
        "DRS_zip", "MANE",
    ]
    for key in required:
        p = PATHS[key]
        ok = p.exists()
        add_qc("asset_exists", key, "PASS" if ok else "FAIL", str(p), "exists")
    if not PATHS["DRS_zip"].exists():
        raise RuntimeError("DRS ZIP missing.")
    try:
        with zipfile.ZipFile(PATHS["DRS_zip"], "r") as zf:
            bad = zf.testzip()
            add_qc("drs_zip_integrity", "DRS", "PASS" if bad is None else "FAIL", bad or "None", "None")
            if bad is not None:
                raise RuntimeError(f"DRS ZIP corrupt member: {bad}")
    except Exception:
        raise

    # Freeze/verify MANE checksum if prior freeze exists.
    freeze = META / "reference_freeze.json"
    if freeze.exists() and PATHS["MANE"].exists():
        try:
            expected = json.loads(freeze.read_text(encoding="utf-8-sig")).get("sha256", "")
            got = sha256(PATHS["MANE"])
            add_qc("mane_sha256", "MANE", "PASS" if (not expected or expected == got) else "FAIL", got, expected)
        except Exception as e:
            add_error("preflight", "MANE", repr(e))


# ---------------------------------------------------------------------
# DRS extraction + exact supplementary table inventory
# ---------------------------------------------------------------------

def extract_drs() -> Dict[str, Path]:
    DRS_EXTRACT.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(PATHS["DRS_zip"], "r") as zf:
        names = zf.namelist()
        for member in names:
            # Protect against path traversal
            target = (DRS_EXTRACT / member).resolve()
            if DRS_EXTRACT.resolve() not in target.parents and target != DRS_EXTRACT.resolve():
                raise RuntimeError(f"Unsafe ZIP member: {member}")
        zf.extractall(DRS_EXTRACT)

    wanted = {}
    for n in ["SupplementaryTableS3.xlsx", "SupplementaryTableS6.xlsx", "SupplementaryTableS8.xlsx"]:
        matches = list(DRS_EXTRACT.rglob(n))
        if matches:
            wanted[n] = matches[0]
            add_qc("drs_table_present", n, "PASS", str(matches[0]), "present")
        else:
            add_qc("drs_table_present", n, "FAIL", "", "present")
    return wanted


# ---------------------------------------------------------------------
# Primary reported-site normalization
# ---------------------------------------------------------------------

def normalize_bid(dataset_id: str, cell_line: str, path: Path) -> List[Dict[str, Any]]:
    recs, cols, _ = read_records(
        path, "Sheet1", 3,
        ["chr", "pos", "strand", "Deletion", "Frac", "Motif"]
    )
    out = []
    for rec in recs:
        chrom = normalize_chrom(rec.get("chr"))
        pos = first_existing(rec, ["pos"])
        strand = normalize_strand(rec.get("strand"))
        m = motif5(first_existing(rec, ["Motif_1", "Motif"]))
        row = {
            "dataset_id": dataset_id,
            "assay": "BID",
            "chemistry_family": CHEMISTRY["BID"],
            "cell_line": cell_line,
            "reported_call": 1,
            "assembly": "GRCh38",
            "chrom": chrom,
            "strand": strand,
            "source_pos": pos,
            "pos1": pos,
            "source_start": "",
            "source_end": "",
            "coordinate_semantics": "1-based genomic nucleotide",
            "base_locus_id": f"{chrom}:{strand}:{safe_str(pos)}" if chrom and strand and safe_str(pos) else "",
            "gene": safe_str(rec.get("name")),
            "transcript_id": safe_str(rec.get("refseq")),
            "region": safe_str(rec.get("seg")),
            "motif_5mer_reported": m,
            "observed_local_u_run_len_5mer": local_urun_len_from_5mer(m),
            "signal_primary_name": "Frac_Ave_percent",
            "signal_primary": first_existing(rec, ["Frac_Ave_%"]),
            "deletion_average": first_existing(rec, ["Deletion_Ave"]),
            "confidence_raw": "",
            "source_file": path.name,
            "source_sheet": "Sheet1",
            "source_row": rec["_excel_row"],
        }
        row.update(extract_raw_columns(rec))
        out.append(row)

    exp = EXPECTED_COUNTS[dataset_id]
    add_qc("row_count", dataset_id, "PASS" if len(out) == exp else "FAIL", len(out), exp)
    if out:
        centered = [r["motif_5mer_reported"] for r in out if r["motif_5mer_reported"]]
        centered_ok = sum(len(m) == 5 and m[2] == "U" for m in centered)
        add_qc(
            "bid_motif_center_U",
            dataset_id,
            "PASS" if centered and centered_ok == len(centered) else "WARN",
            f"{centered_ok}/{len(centered)}",
            "all usable motifs centered U",
        )
    return out


def normalize_elap(dataset_id: str, cell_line: str, path: Path) -> List[Dict[str, Any]]:
    recs, cols, _ = read_records(
        path, "Sheet1", 0,
        ["chrom", "start", "end", "strand", "ref", "confidence", "enrich"]
    )
    out = []
    for rec in recs:
        chrom = normalize_chrom(rec.get("chrom"))
        start = rec.get("start")
        end = rec.get("end")
        strand = normalize_strand(rec.get("strand"))
        # Validated Stage0 convention: one-base [start,end), pos1=end.
        pos1 = end
        row = {
            "dataset_id": dataset_id,
            "assay": "ELAP",
            "chemistry_family": CHEMISTRY["ELAP"],
            "cell_line": cell_line,
            "reported_call": 1,
            "assembly": "GRCh38",
            "chrom": chrom,
            "strand": strand,
            "source_pos": "",
            "pos1": pos1,
            "source_start": start,
            "source_end": end,
            "coordinate_semantics": "one-base BED-like [start,end); normalized pos1=end",
            "base_locus_id": f"{chrom}:{strand}:{safe_str(pos1)}" if chrom and strand and safe_str(pos1) else "",
            "gene": "",
            "transcript_id": "",
            "region": "",
            "motif_5mer_reported": "",
            "observed_local_u_run_len_5mer": "",
            "signal_primary_name": "avg(stop_ratio*stopped_reads)",
            "signal_primary": first_existing(rec, ["avg(stop_ratio*stopped_reads)"]),
            "enrich_rep1": first_existing(rec, ["enrich_rep1"]),
            "enrich_rep2": first_existing(rec, ["enrich_rep2"]),
            "confidence_raw": first_existing(rec, ["confidence"]),
            "modification_estimation": first_existing(rec, ["modification_estimation"]),
            "source_file": path.name,
            "source_sheet": "Sheet1",
            "source_row": rec["_excel_row"],
        }
        row.update(extract_raw_columns(rec))
        out.append(row)

    exp = EXPECTED_COUNTS[dataset_id]
    add_qc("row_count", dataset_id, "PASS" if len(out) == exp else "FAIL", len(out), exp)
    onebase = 0
    for r in out:
        a = to_float(r["source_start"])
        b = to_float(r["source_end"])
        if a is not None and b is not None and b - a == 1:
            onebase += 1
    add_qc(
        "elap_one_base_intervals", dataset_id,
        "PASS" if out and onebase == len(out) else "FAIL",
        f"{onebase}/{len(out)}", "all [start,end) length 1"
    )
    return out


def normalize_bacs(path: Path) -> List[Dict[str, Any]]:
    sheet = "Supplementary Table 8"
    recs, cols, _ = read_records(
        path, sheet, 2,
        ["chr", "pos", "strand", "motif", "gene", "BACS", "Ψ level", "transcript"]
    )
    out = []
    for rec in recs:
        chrom = normalize_chrom(rec.get("chr"))
        pos = rec.get("pos")
        strand = normalize_strand(rec.get("strand"))
        m = motif5(rec.get("motif"))
        row = {
            "dataset_id": "BACS_HeLa",
            "assay": "BACS",
            "chemistry_family": CHEMISTRY["BACS"],
            "cell_line": "HeLa",
            "reported_call": 1,
            "assembly": "GRCh38",
            "chrom": chrom,
            "strand": strand,
            "source_pos": pos,
            "pos1": pos,
            "source_start": "",
            "source_end": "",
            "coordinate_semantics": "1-based genomic nucleotide",
            "base_locus_id": f"{chrom}:{strand}:{safe_str(pos)}" if chrom and strand and safe_str(pos) else "",
            "gene": safe_str(rec.get("gene")),
            "transcript_id": safe_str(rec.get("transcript_id")),
            "region": safe_str(rec.get("feature")),
            "motif_5mer_reported": m,
            "observed_local_u_run_len_5mer": local_urun_len_from_5mer(m),
            "signal_primary_name": "psi_level",
            "signal_primary": first_existing(rec, ["psi_level", "Ψ_level"]),
            "bacs_conversion_rate": first_existing(rec, ["BACS_conversion_rate"]),
            "control_conversion_rate": first_existing(rec, ["control_conversion_rate"]),
            "ivt_conversion_rate": first_existing(rec, ["IVT_conversion_rate"]),
            "confidence_raw": first_existing(rec, ["p-value", "p_value"]),
            "source_file": path.name,
            "source_sheet": sheet,
            "source_row": rec["_excel_row"],
        }
        row.update(extract_raw_columns(rec))
        out.append(row)

    exp = EXPECTED_COUNTS["BACS_HeLa"]
    add_qc("row_count", "BACS_HeLa", "PASS" if len(out) == exp else "FAIL", len(out), exp)
    motifs = [r["motif_5mer_reported"] for r in out if r["motif_5mer_reported"]]
    centered = sum(m[2] == "U" for m in motifs)
    add_qc("bacs_motif_center_U", "BACS_HeLa",
           "PASS" if motifs and centered == len(motifs) else "FAIL",
           f"{centered}/{len(motifs)}", "all usable motifs centered U")
    return out


def parse_praise_chr_site(s: str) -> Tuple[str, str, str, str]:
    """
    Preserve PRAISE semantics. We parse the published string but do not force a unique nucleotide.
    Returns chrom, token1, token2, interval_flag.
    """
    s = safe_str(s)
    m = re.match(r"^(chr[^_]+)_(\d+)(?:-(\d+))?$", s)
    if not m:
        return "", "", "", ""
    chrom, a, b = m.group(1), m.group(2), m.group(3)
    return chrom, a, (b or ""), "1" if b else "0"


def normalize_praise(path: Path) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    # Dataset 2 = reported quantitative HEK293T sites
    matrix = read_matrix(path, "Supplementary Dataset 2")
    h = detect_header_row(matrix, ["chr", "site", "gene", "transcript", "rep"])
    recs, cols, mapping = records_from_matrix(matrix, h)
    for cleaned, original in mapping:
        column_dictionary_rows.append({
            "source_file": path.name,
            "source_sheet": "Supplementary Dataset 2",
            "clean_column": cleaned,
            "original_column": original,
        })
    schema_rows.append({
        "source_file": path.name,
        "source_sheet": "Supplementary Dataset 2",
        "header_row_1based": h + 1,
        "data_rows": len(recs),
        "n_columns": len(cols),
        "columns": " | ".join(cols),
    })

    # Find likely columns by semantic names, but preserve everything.
    def find_col(tokens: Sequence[str]) -> Optional[str]:
        scored = []
        for c in cols:
            cl = c.lower()
            score = sum(t.lower() in cl for t in tokens)
            if score:
                scored.append((score, c))
        return sorted(scored, reverse=True)[0][1] if scored else None

    chr_site_col = find_col(["chr_site", "chrsite", "genomic"])
    # fallback: detect values like chr14_732...
    if chr_site_col is None:
        for c in cols:
            vals = [safe_str(r.get(c)) for r in recs[:100]]
            if sum(bool(re.match(r"^chr[^_]+_\d+", v)) for v in vals) >= 10:
                chr_site_col = c
                break
    gene_col = find_col(["gene"])
    tx_col = find_col(["transcript", "refseq"])
    site_col = find_col(["site"])
    signal_cols = [c for c in cols if any(t in c.lower() for t in ["rep", "level", "ratio", "fraction"])]

    out = []
    for rec in recs:
        chr_site = safe_str(rec.get(chr_site_col)) if chr_site_col else ""
        chrom, a, b, interval_flag = parse_praise_chr_site(chr_site)
        row = {
            "dataset_id": "PRAISE_HEK293T",
            "assay": "PRAISE",
            "chemistry_family": CHEMISTRY["PRAISE"],
            "cell_line": "HEK293T",
            "reported_call": 1,
            "assembly": "GRCh38",
            "chrom": chrom,
            "strand": "",
            "source_pos": safe_str(rec.get(site_col)) if site_col else "",
            "pos1": "",
            "source_start": a,
            "source_end": b,
            "coordinate_semantics": "published chr_site nucleotide/interval; no forced exact base in Phase0C",
            "base_locus_id": "",
            "gene": safe_str(rec.get(gene_col)) if gene_col else "",
            "transcript_id": safe_str(rec.get(tx_col)) if tx_col else "",
            "region": "",
            "motif_5mer_reported": "",
            "observed_local_u_run_len_5mer": "",
            "signal_primary_name": signal_cols[0] if signal_cols else "",
            "signal_primary": rec.get(signal_cols[0]) if signal_cols else "",
            "confidence_raw": "",
            "praise_chr_site": chr_site,
            "praise_interval_flag": interval_flag,
            "source_file": path.name,
            "source_sheet": "Supplementary Dataset 2",
            "source_row": rec["_excel_row"],
        }
        row.update(extract_raw_columns(rec))
        out.append(row)

    add_qc("row_count", "PRAISE_HEK293T",
           "PASS" if len(out) == EXPECTED_COUNTS["PRAISE_HEK293T"] else "FAIL",
           len(out), EXPECTED_COUNTS["PRAISE_HEK293T"])
    intervals = sum(r["praise_interval_flag"] == "1" for r in out)
    add_qc("praise_interval_records", "PRAISE_HEK293T", "PASS",
           intervals, "historical Stage0 value around 1357",
           "Diagnostic only; exact semantics remain chemistry-aware.")

    # Dataset 3 = writer assignment / perturbation resource
    matrix3 = read_matrix(path, "Supplementary Dataset 3")
    h3 = detect_header_row(matrix3, ["gene", "PUS", "TRUB1", "DKC1", "site"])
    recs3, cols3, mapping3 = records_from_matrix(matrix3, h3)
    for cleaned, original in mapping3:
        column_dictionary_rows.append({
            "source_file": path.name,
            "source_sheet": "Supplementary Dataset 3",
            "clean_column": cleaned,
            "original_column": original,
        })
    schema_rows.append({
        "source_file": path.name,
        "source_sheet": "Supplementary Dataset 3",
        "header_row_1based": h3 + 1,
        "data_rows": len(recs3),
        "n_columns": len(cols3),
        "columns": " | ".join(cols3),
    })
    writer_rows = []
    for rec in recs3:
        rr = {
            "source_file": path.name,
            "source_sheet": "Supplementary Dataset 3",
            "source_row": rec["_excel_row"],
        }
        rr.update({f"raw__{k}": v for k, v in rec.items() if k != "_excel_row"})
        writer_rows.append(rr)

    return out, writer_rows


# ---------------------------------------------------------------------
# DRS S3/S6/S8 normalization: provenance-first, no forced coordinates yet
# ---------------------------------------------------------------------

def normalize_drs_table(path: Path, table_id: str) -> List[Dict[str, Any]]:
    out = []
    for sheet in list_sheets(path):
        matrix = read_matrix(path, sheet)
        h = detect_header_row(
            matrix,
            ["gene", "transcript", "chrom", "position", "HeLa", "A549", "TPM", "protein", "TE", "psi", "pseudouridine"]
        )
        recs, cols, mapping = records_from_matrix(matrix, h)
        for cleaned, original in mapping:
            column_dictionary_rows.append({
                "source_file": path.name,
                "source_sheet": sheet,
                "clean_column": cleaned,
                "original_column": original,
            })
        schema_rows.append({
            "source_file": path.name,
            "source_sheet": sheet,
            "header_row_1based": h + 1,
            "data_rows": len(recs),
            "n_columns": len(cols),
            "columns": " | ".join(cols),
        })
        for rec in recs:
            rr = {
                "drs_table": table_id,
                "assay": "DRS",
                "chemistry_family": CHEMISTRY["DRS"],
                "source_file": path.name,
                "source_sheet": sheet,
                "source_row": rec["_excel_row"],
            }
            rr.update(extract_raw_columns(rec))
            out.append(rr)
    return out


# ---------------------------------------------------------------------
# Synthetic calibration extraction
# ---------------------------------------------------------------------

def motif_like(x: Any) -> str:
    s = safe_str(x).upper().replace("Ψ", "U").replace("T", "U")
    s = s.replace(" ", "").replace("-", "")
    # exact 5-mer only, central U
    if len(s) == 5 and re.fullmatch(r"[ACGU]{5}", s) and s[2] == "U":
        return s
    return ""


def detect_calibration_sheet(path: Path) -> Optional[Dict[str, Any]]:
    best = None
    for sheet in list_sheets(path):
        matrix = read_matrix(path, sheet)
        if not matrix:
            continue
        max_cols = max(len(r) for r in matrix)
        for c in range(max_cols):
            motifs = []
            rows = []
            for ri, row in enumerate(matrix):
                v = row[c] if c < len(row) else None
                m = motif_like(v)
                if m:
                    motifs.append(m)
                    rows.append(ri)
            uniq = len(set(motifs))
            if uniq >= 100:
                cand = {
                    "sheet": sheet,
                    "motif_col_idx": c,
                    "motifs": motifs,
                    "row_indices": rows,
                    "unique_motifs": uniq,
                    "matrix": matrix,
                }
                if best is None or uniq > best["unique_motifs"]:
                    best = cand
    return best


def header_for_calibration(matrix: Sequence[Sequence[Any]], first_data_row: int) -> Tuple[int, List[str]]:
    # search up to 8 rows above first motif row; prefer the closest text-rich row
    lo = max(0, first_data_row - 8)
    candidates = []
    for i in range(lo, first_data_row):
        row = matrix[i]
        vals = [safe_str(v) for v in row]
        nonempty = sum(bool(v) for v in vals)
        textual = sum(bool(v) and not re.fullmatch(r"[-+]?\d+(\.\d+)?", v) for v in vals)
        candidates.append((textual * 2 + nonempty, i, vals))
    if not candidates:
        return max(0, first_data_row - 1), []
    _, i, vals = sorted(candidates, reverse=True)[0]
    cols, _ = make_unique_columns(vals)
    return i, cols


def score_calibration_column(assay: str, header: str) -> int:
    h = header.lower()
    if assay == "BID":
        if h == "r":
            return 100
        score = 0
        if "deletion" in h:
            score += 20
        if "100" in h or "psi" in h or "modified" in h or "induced" in h:
            score += 10
        if "background" in h or h == "b" or "unmodified" in h or "input" in h:
            score -= 30
        return score
    if assay == "BACS":
        score = 0
        if "conversion" in h:
            score += 25
        if "psi" in h or "nnpsinn" in h or "nnunn" not in h:
            score += 5
        if "control" in h or "unmodified" in h or "false" in h or "fp" in h:
            score -= 30
        return score
    if assay == "ELAP":
        score = 0
        if "enrich" in h:
            score += 30
        if "relative" in h or "normalized" in h or "ratio" in h or "fold" in h:
            score += 15
        if "read" in h or "count" in h or "abundance" in h:
            score += 8
        if "control" in h or "input" in h or "unmodified" in h:
            score -= 20
        return score
    return 0


def extract_calibration(assay: str, path: Path) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    det = detect_calibration_sheet(path)
    if det is None:
        return [], {
            "assay": assay, "source_file": path.name, "status": "NO_5MER_BLOCK",
            "sheet": "", "unique_motifs": 0, "selected_response": "", "notes": ""
        }

    matrix = det["matrix"]
    rows = det["row_indices"]
    motif_col = det["motif_col_idx"]
    first = min(rows)
    header_idx, headers = header_for_calibration(matrix, first)

    max_cols = max(len(r) for r in matrix)
    if len(headers) < max_cols:
        headers = headers + [f"col_{i+1}" for i in range(len(headers), max_cols)]

    # Evaluate numeric coverage + scientific semantic score.
    numeric_candidates = []
    for c in range(max_cols):
        if c == motif_col:
            continue
        vals = []
        for ri in rows:
            row = matrix[ri]
            v = to_float(row[c] if c < len(row) else None)
            if v is not None:
                vals.append(v)
        if len(vals) >= max(50, int(0.5 * len(rows))):
            header = headers[c] if c < len(headers) else f"col_{c+1}"
            semantic = score_calibration_column(assay, header)
            numeric_candidates.append((semantic, len(vals), c, header))

    numeric_candidates.sort(reverse=True)
    selected = None
    if numeric_candidates and numeric_candidates[0][0] > 0:
        selected = numeric_candidates[0]

    raw_rows = []
    for ri in rows:
        row = matrix[ri]
        m = motif_like(row[motif_col] if motif_col < len(row) else None)
        if not m:
            continue
        rr = {
            "assay": assay,
            "chemistry_family": CHEMISTRY[assay],
            "motif_RNA5": m,
            "source_file": path.name,
            "source_sheet": det["sheet"],
            "source_row": ri + 1,
        }
        for c in range(max_cols):
            val = row[c] if c < len(row) else None
            h = headers[c] if c < len(headers) else f"col_{c+1}"
            rr[f"raw__{clean_col(h)}"] = val
        raw_rows.append(rr)

    selected_name = ""
    if selected:
        _, _, sel_c, sel_h = selected
        selected_name = sel_h
        values = []
        for ri in rows:
            row = matrix[ri]
            values.append(to_float(row[sel_c] if sel_c < len(row) else None))
        ranks = percentile_rank(values)

        # attach selected response by source row order
        rank_map = {}
        val_map = {}
        for ri, v, pr in zip(rows, values, ranks):
            rank_map[ri + 1] = pr
            val_map[ri + 1] = v
        for rr in raw_rows:
            rr["calibration_response_name"] = selected_name
            rr["calibration_response_raw"] = val_map.get(rr["source_row"])
            rr["calibration_response_percentile"] = rank_map.get(rr["source_row"])

    mapping = {
        "assay": assay,
        "source_file": path.name,
        "status": "SELECTED" if selected else "RESPONSE_UNRESOLVED",
        "sheet": det["sheet"],
        "header_row_1based": header_idx + 1,
        "motif_column_index_1based": motif_col + 1,
        "unique_motifs": len(set(r["motif_RNA5"] for r in raw_rows)),
        "selected_response": selected_name,
        "candidate_numeric_columns": " | ".join(
            f"{h}[semantic={s},n={n}]" for s, n, c, h in numeric_candidates[:12]
        ),
        "notes": "Response selection is based on assay chemistry/header semantics only, never downstream performance.",
    }
    return raw_rows, mapping


# ---------------------------------------------------------------------
# Main execution
# ---------------------------------------------------------------------

def main() -> int:
    print("=" * 94)
    print("TRACE PHASE 0C — FORMAL SOURCE NORMALIZATION")
    print("=" * 94)

    preflight()

    # DRS extraction
    drs_paths = extract_drs()

    # Primary site maps
    normalized_tables: Dict[str, List[Dict[str, Any]]] = {}
    normalized_tables["BID_HEK293T"] = normalize_bid("BID_HEK293T", "HEK293T", PATHS["BID_HEK293T"])
    normalized_tables["BID_HeLa"] = normalize_bid("BID_HeLa", "HeLa", PATHS["BID_HeLa"])
    normalized_tables["BID_A549"] = normalize_bid("BID_A549", "A549", PATHS["BID_A549"])
    normalized_tables["ELAP_HEK293T"] = normalize_elap("ELAP_HEK293T", "HEK293T", PATHS["ELAP_HEK293T"])
    normalized_tables["ELAP_HeLa"] = normalize_elap("ELAP_HeLa", "HeLa", PATHS["ELAP_HeLa"])
    normalized_tables["BACS_HeLa"] = normalize_bacs(PATHS["BACS_HeLa"])
    praise_sites, praise_writers = normalize_praise(PATHS["PRAISE"])
    normalized_tables["PRAISE_HEK293T"] = praise_sites

    # Save normalized site tables
    for dataset_id, rows in normalized_tables.items():
        out = NORM / f"{dataset_id}.tsv"
        write_tsv(out, rows)
        manifest_rows.append({
            "dataset_id": dataset_id,
            "kind": "reported_site_table",
            "output_path": str(out),
            "n_rows": len(rows),
            "sha256": sha256(out),
            "status": "NORMALIZED",
        })

    writer_out = NORM / "PRAISE_writer_assignments.tsv"
    write_tsv(writer_out, praise_writers)
    manifest_rows.append({
        "dataset_id": "PRAISE_writer_assignments",
        "kind": "writer_validation",
        "output_path": str(writer_out),
        "n_rows": len(praise_writers),
        "sha256": sha256(writer_out),
        "status": "NORMALIZED",
    })

    # DRS S3/S6/S8: preserve every published field, no forced site semantics yet.
    for fname, table_id in [
        ("SupplementaryTableS3.xlsx", "S3_functional"),
        ("SupplementaryTableS6.xlsx", "S6_psi_calls"),
        ("SupplementaryTableS8.xlsx", "S8_conserved_controls"),
    ]:
        p = drs_paths.get(fname)
        if not p:
            continue
        rows = normalize_drs_table(p, table_id)
        out = NORM / f"DRS_{table_id}.tsv"
        write_tsv(out, rows)
        manifest_rows.append({
            "dataset_id": f"DRS_{table_id}",
            "kind": "heldout_or_functional_source",
            "output_path": str(out),
            "n_rows": len(rows),
            "sha256": sha256(out),
            "status": "NORMALIZED_PROVENANCE_ONLY",
        })

    # Synthetic calibration.
    # BID: use Supplementary Table workbook (A/B/R calibration parameters).
    # BACS: use Fig.1 source data.
    # ELAP: use Source Data workbook (Fig.2 contains 256-context sequence preference).
    calibration_all = []
    for assay, p in [
        ("BID", PATHS["BID_supp"]),
        ("BACS", PATHS["BACS_fig1"]),
        ("ELAP", PATHS["ELAP_source"]),
    ]:
        try:
            rows, mapping = extract_calibration(assay, p)
            calibration_all.extend(rows)
            calibration_mapping_rows.append(mapping)
            nuniq = mapping.get("unique_motifs", 0)
            status = "PASS" if nuniq >= 200 and mapping.get("status") == "SELECTED" else "REVIEW"
            add_qc("synthetic_calibration", assay, status, nuniq, ">=200 unique NNUNN contexts",
                   mapping.get("selected_response", ""))
        except Exception as e:
            add_error("calibration", assay, traceback.format_exc())
            calibration_mapping_rows.append({
                "assay": assay,
                "source_file": p.name,
                "status": "ERROR",
                "sheet": "",
                "unique_motifs": 0,
                "selected_response": "",
                "notes": repr(e),
            })
            add_qc("synthetic_calibration", assay, "FAIL", 0, ">=200 unique contexts", repr(e))

    cal_out = HARM / "synthetic_context_calibration.tsv"
    write_tsv(cal_out, calibration_all)
    manifest_rows.append({
        "dataset_id": "synthetic_context_calibration",
        "kind": "calibration",
        "output_path": str(cal_out),
        "n_rows": len(calibration_all),
        "sha256": sha256(cal_out),
        "status": "NORMALIZED_WITH_MAPPING_REPORT",
    })

    # Data contract
    contract = {
        "project": "Trace",
        "phase": "0C",
        "scientific_story": "Calibration -> Measurement -> Evidence -> Inference",
        "estimand_boundary": (
            "Published reported sites are observations conditioned on assay/calling pipeline. "
            "Absence from another published call table is NOT a biological negative."
        ),
        "chemistry_families": CHEMISTRY,
        "coordinate_contract": COORDINATE_CONTRACT,
        "primary_site_tables": {
            k: {
                "expected_rows": EXPECTED_COUNTS[k],
                "normalized_file": str(NORM / f"{k}.tsv"),
            }
            for k in EXPECTED_COUNTS
        },
        "aim1_primary_pair": {
            "cells": ["HEK293T", "HeLa"],
            "assays": ["BID", "ELAP"],
            "training_labels_not_created_in_phase0c": True,
        },
        "aim1_second_platform": {
            "matched_cells": ["HeLa", "A549"],
            "assays": ["BID", "DRS"],
            "DRS_is_ground_truth": False,
        },
        "synthetic_calibration_policy": {
            "compare_raw_units_across_assays": False,
            "within_assay_percentile_normalization": True,
            "response_selection_basis": "published chemistry / source-data semantics, never downstream model performance",
        },
        "locus_policy": {
            "base_locus_id": "exact nucleotide only when published coordinate semantics support it",
            "equivalence_locus_id": "deferred until chemistry-aware U-run reconstruction",
            "arbitrary_plusminus_k_matching": False,
        },
    }
    (META / "phase0c_data_contract.json").write_text(
        json.dumps(contract, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    # Write metadata outputs
    write_tsv(META / "phase0c_qc.tsv", qc_rows)
    write_tsv(META / "phase0c_errors.tsv", error_rows)
    write_tsv(META / "phase0c_schema.tsv", schema_rows)
    write_tsv(META / "phase0c_normalized_manifest.tsv", manifest_rows)
    write_tsv(META / "phase0c_calibration_mapping.tsv", calibration_mapping_rows)
    write_tsv(META / "phase0c_column_dictionary.tsv", column_dictionary_rows)

    required_qc_fail = [
        r for r in qc_rows
        if r["status"] == "FAIL"
        and r["check_id"] in {
            "asset_exists", "drs_zip_integrity", "mane_sha256",
            "row_count", "elap_one_base_intervals", "bacs_motif_center_U",
            "drs_table_present"
        }
    ]
    cal_review = [
        r for r in qc_rows
        if r["check_id"] == "synthetic_calibration" and r["status"] != "PASS"
    ]

    if required_qc_fail:
        gate = "STOP_SOURCE_NORMALIZATION_QC_FAILED"
    elif cal_review:
        gate = "SITE_NORMALIZATION_PASS_CALIBRATION_MAPPING_REVIEW"
    else:
        gate = "READY_FOR_CANONICAL_HARMONIZATION_AND_AIM1A"

    summary = {
        "phase": "0C",
        "status": gate,
        "normalized_site_datasets": {k: len(v) for k, v in normalized_tables.items()},
        "praise_writer_rows": len(praise_writers),
        "calibration_mapping": calibration_mapping_rows,
        "qc_pass": sum(r["status"] == "PASS" for r in qc_rows),
        "qc_review_or_warn": sum(r["status"] in {"REVIEW", "WARN"} for r in qc_rows),
        "qc_fail": sum(r["status"] == "FAIL" for r in qc_rows),
        "errors": len(error_rows),
        "next_if_ready": (
            "Build chemistry-aware canonical loci, add a single frozen GRCh38 FASTA for common sequence extraction, "
            "then run Aim1A synthetic-calibration transfer and Aim1B HEK<->HeLa cross-cell transfer."
        ),
    }
    (META / "phase0c_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    print()
    print("=" * 94)
    print("PHASE 0C COMPLETE")
    print("=" * 94)
    print("STATUS:", gate)
    print()
    for k, v in summary["normalized_site_datasets"].items():
        print(f"{k:22s} {v:6d}")
    print(f"{'PRAISE writers':22s} {len(praise_writers):6d}")
    print()
    print("Calibration:")
    for m in calibration_mapping_rows:
        print(
            f"  {m.get('assay',''):5s} "
            f"status={m.get('status','')} "
            f"motifs={m.get('unique_motifs',0)} "
            f"response={m.get('selected_response','')}"
        )
    print()
    print(f"QC FAIL: {summary['qc_fail']}   Errors: {summary['errors']}")
    print()
    print(r"Return/upload these files from D:\RNA\Trace:")
    print(r"  00_meta\phase0c_summary.json")
    print(r"  00_meta\phase0c_qc.tsv")
    print(r"  00_meta\phase0c_calibration_mapping.tsv")
    print(r"  00_meta\phase0c_schema.tsv")
    print(r"  00_meta\phase0c_normalized_manifest.tsv")
    print(r"  00_meta\phase0c_errors.tsv")
    print(r"  00_meta\phase0c_data_contract.json")
    print(r"  03_harmonized\synthetic_context_calibration.tsv")
    print()
    print("Do NOT run any classifier yet.")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception:
        err = traceback.format_exc()
        try:
            LOGS.mkdir(parents=True, exist_ok=True)
            (LOGS / "phase0c_FATAL_ERROR.txt").write_text(err, encoding="utf-8")
        except Exception:
            pass
        print("\nFATAL ERROR\n")
        print(err)
        print(r"Send D:\RNA\Trace\logs\phase0c_FATAL_ERROR.txt back to ChatGPT.")
        try:
            input("\nPress Enter to close...")
        except Exception:
            pass
        raise
