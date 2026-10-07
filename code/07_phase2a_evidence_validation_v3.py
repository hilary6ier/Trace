#!/usr/bin/env python3
# -*- coding: utf-8 -*-
r"""
TRACE Phase 2A v3 — chemistry-diverse support and independent biological evidence
==============================================================================

Scientific question
-------------------
Among Ψ sites already reported by one source assay, does support from
independent biochemical measurement principles carry additional biological
evidence?

Primary source-positive universe
--------------------------------
BID-seq HeLa shControl mRNA Ψ sites used in the original PUS-knockdown
experiment (published n=133, all >10% Ψ fraction).

Outcome
-------
"writer_assigned" means the published site was reported as responsive to at
least one of the eight tested PUS knockdowns in the BID-seq experiment.
This is NOT "true Ψ" and the complementary group is NOT a biological negative.

External evidence
-----------------
BACS HeLa  = cyclization / U->C conversion
ELAP HeLa  = enzymatic labeling / RT-stop
DRS HeLa   = native direct-RNA platform, HELD OUT from evidence construction

Primary evidence variable
-------------------------
chemistry_breadth = BACS support + ELAP support (0,1,2)

DRS is never included in chemistry_breadth. It is used only afterwards as a
held-out reporting-support test.

Matching
--------
For BID/BACS/ELAP:
- exact nucleotide match always counts;
- same contiguous-U run counts only when the full local U-run is recoverable
  from the published 5-mer in BOTH datasets;
- edge-truncated 5-mer runs fall back to exact matching.
No arbitrary +/-k window is used.

Statistics
----------
1) PUS-assignment rate by chemistry breadth with Wilson intervals.
2) Primary logistic model:
       writer_assigned ~ chemistry_breadth + source_fraction_rank
   with motif-cluster robust SE.
3) Sensitivity:
       + known TRUB1/PUS7 motif indicator
       exact-coordinate-only evidence
       locally single-U source sites
4) Held-out DRS:
       DRS reported support ~ chemistry_breadth + source_fraction_rank
   and descriptive support gradient. DRS absence is interpreted only as
   "not reported by the DRS published calling pipeline", never as unmodified.

Hard provenance gates
---------------------
- shControl source table must contain 133 sites.
- BID Supplementary Table 8 must yield 133 source sites. The author-defined 104-site outcome is recovered from the separate official Source Data Fig.3k heatmap membership. No numeric PUS-response threshold is inferred.
- External BACS/ELAP tables must be the already normalized Trace assets.

No FASTQ/BAM/raw sequencing.
"""

from __future__ import annotations

import json
import math
import re
import traceback
import urllib.request
import zipfile
from collections import defaultdict
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Set, Tuple

import numpy as np
import pandas as pd
import openpyxl

ROOT = Path(r"D:\RNA\Trace")
META = ROOT / "00_meta"
SRC = ROOT / "01_source"
NORM = ROOT / "03_harmonized" / "normalized_sources"
AIM2 = ROOT / "05_aim2"
LOGS = ROOT / "logs"

for p in [META, AIM2, LOGS]:
    p.mkdir(parents=True, exist_ok=True)

BID_SH = SRC / "BID" / "GSE179798_HeLa_mRNA_shControl_BID-seq.xlsx"
BID_SUPP = SRC / "BID" / "BID_supplementary_tables.xlsx"
BID_FIG3 = SRC / "BID" / "BID_Source_Data_Fig3.xlsx"
BID_FIG3_URL = (
    "https://media.springernature.com/original/springer-static/esm/"
    "art%3A10.1038%2Fs41587-022-01505-w/MediaObjects/"
    "41587_2022_1505_MOESM16_ESM.xlsx"
)
BACS = NORM / "BACS_HeLa.tsv"
ELAP = NORM / "ELAP_HeLa.tsv"
DRS = NORM / "DRS_S6_psi_calls.tsv"
BACS_BOOK = SRC / "BACS" / "GSE241849_Supplementary_Table_3_10.xlsx"

SEED = 20261001
RNG = np.random.default_rng(SEED)

WRITERS = ["TRUB1", "PUS7", "PUS1", "PUS3", "PUS7L", "PUSL1", "TRUB2", "DKC1"]
PUS7_RE = re.compile(r"^U[ACGU]UA[AG]$")
TRUB1_RE = re.compile(r"^GUUC[ACGU]$")


# -----------------------------------------------------------------------------
# generic helpers
# -----------------------------------------------------------------------------

def ss(x: Any) -> str:
    if x is None:
        return ""
    if isinstance(x, float) and np.isnan(x):
        return ""
    return str(x).strip()


def clean_col(x: Any) -> str:
    s = ss(x).replace("Ψ", "psi")
    s = re.sub(r"\s+", "_", s)
    s = re.sub(r"[^A-Za-z0-9_.()+%/-]+", "_", s)
    s = re.sub(r"_+", "_", s).strip("_")
    return s or "unnamed"


def motif5(x: Any) -> str:
    z = ss(x).upper().replace("Ψ", "U").replace("T", "U")
    z = re.sub(r"[^ACGU]", "", z)
    return z if len(z) == 5 and z[2] == "U" else ""


def norm_chr(x: Any) -> str:
    z = ss(x)
    if not z:
        return ""
    if z.lower().startswith("chr"):
        return "chr" + z[3:]
    if re.fullmatch(r"(?:\d+|X|Y|M|MT)", z, flags=re.I):
        return "chr" + ("M" if z.upper() == "MT" else z)
    return z


def norm_strand(x: Any) -> str:
    z = ss(x)
    return z if z in {"+", "-"} else ""


def as_int(x: Any) -> Optional[int]:
    try:
        return int(float(x))
    except Exception:
        return None


def as_float(x: Any) -> Optional[float]:
    try:
        v = float(str(x).replace("%", "").replace(",", ""))
        return v if np.isfinite(v) else None
    except Exception:
        return None


def percentile_rank(x: pd.Series) -> pd.Series:
    return pd.to_numeric(x, errors="coerce").rank(method="average", pct=True)


def wilson(k: int, n: int, z: float = 1.959963984540054) -> Tuple[float, float]:
    if n == 0:
        return np.nan, np.nan
    p = k / n
    den = 1 + z*z/n
    ctr = (p + z*z/(2*n)) / den
    rad = z * math.sqrt(p*(1-p)/n + z*z/(4*n*n)) / den
    return ctr-rad, ctr+rad


# -----------------------------------------------------------------------------
# read / normalize BID shControl
# -----------------------------------------------------------------------------

def read_bid_shcontrol() -> pd.DataFrame:
    # Same published BID workbook structure as WT assets; header is row 4.
    raw = pd.read_excel(BID_SH, sheet_name="Sheet1", header=3)
    raw.columns = [clean_col(c) for c in raw.columns]

    required = {"chr", "pos", "strand"}
    if not required.issubset(raw.columns):
        raise RuntimeError(
            f"BID shControl schema unexpected. Missing {sorted(required-set(raw.columns))}; "
            f"columns={list(raw.columns)}"
        )

    motif_col = next((c for c in ["Motif_1", "Motif", "motif"] if c in raw.columns), None)
    frac_col = next(
        (c for c in raw.columns if c.lower() in {"frac_ave_%", "frac_ave", "fraction_ave_%"}),
        None,
    )
    if motif_col is None or frac_col is None:
        raise RuntimeError(
            f"BID shControl motif/fraction columns unresolved. columns={list(raw.columns)}"
        )

    x = pd.DataFrame({
        "chrom": raw["chr"].map(norm_chr),
        "pos1": raw["pos"].map(as_int),
        "strand": raw["strand"].map(norm_strand),
        "motif": raw[motif_col].map(motif5),
        "source_fraction": pd.to_numeric(raw[frac_col], errors="coerce"),
        "source_row": np.arange(len(raw)) + 5,
    })
    x["base_locus_id"] = (
        x["chrom"] + ":" + x["strand"] + ":" + x["pos1"].astype("Int64").astype(str)
    )

    x = x[x["chrom"].ne("") & x["pos1"].notna() & x["strand"].ne("")].copy()

    if len(x) != 133:
        raise RuntimeError(f"BID shControl expected 133 mRNA sites, observed {len(x)}.")
    if x["motif"].eq("").any():
        raise RuntimeError("BID shControl contains unresolved 5-mer motif(s).")
    if x["base_locus_id"].duplicated().any():
        raise RuntimeError("BID shControl contains duplicate exact loci.")

    x["source_fraction_rank"] = percentile_rank(x["source_fraction"])
    return x


# -----------------------------------------------------------------------------
# published BID Table 8 writer assignment
# -----------------------------------------------------------------------------

def detect_header(matrix: List[List[Any]]) -> int:
    tokens = ["chr", "pos", "strand", "gene", "motif", "pus", "trub", "dkc"]
    best = (-1, 0)
    for i, row in enumerate(matrix[:25]):
        text = " ".join(ss(v).lower() for v in row)
        nonempty = sum(bool(ss(v)) for v in row)
        hits = sum(t in text for t in tokens)
        score = hits * 10 + nonempty
        if score > best[0]:
            best = (score, i)
    return best[1]


def find_table8_sheet(path: Path) -> str:
    xl = pd.ExcelFile(path)
    exact = [
        s for s in xl.sheet_names
        if re.search(r"(?i)supplementary\s*table\s*8(?:\D|$)", s)
        or re.fullmatch(r"(?i).*table\s*8", s.strip())
    ]
    if exact:
        return exact[0]
    # conservative fallback: a sheet with token 8 but not 18.
    cands = [s for s in xl.sheet_names if re.search(r"(?<!\d)8(?!\d)", s)]
    if len(cands) == 1:
        return cands[0]
    raise RuntimeError(f"Could not uniquely identify BID Supplementary Table 8: {xl.sheet_names}")


PUBLISHED_WRITER_COUNTS = {
    "TRUB1": 70,
    "PUS7": 40,
    "PUS1": 28,
    "PUS3": 30,
    "PUS7L": 24,
    "PUSL1": 28,
    "TRUB2": 28,
    "DKC1": 33,
}
PUBLISHED_WRITER_UNION = 104


def ensure_bid_fig3_source() -> Path:
    """
    Source Data Fig.3 is a separate official Nature XLSX.
    Supplementary Table 8 is only the 133-site shControl baseline table.
    """
    if BID_FIG3.exists() and zipfile.is_zipfile(BID_FIG3):
        return BID_FIG3

    tmp = BID_FIG3.with_suffix(".download.tmp")
    req = urllib.request.Request(
        BID_FIG3_URL,
        headers={
            "User-Agent": "Mozilla/5.0",
            "Accept": (
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,"
                "application/octet-stream,*/*"
            ),
        },
    )
    try:
        with urllib.request.urlopen(req, timeout=60) as resp, tmp.open("wb") as fout:
            while True:
                block = resp.read(1024 * 1024)
                if not block:
                    break
                fout.write(block)

        if not zipfile.is_zipfile(tmp):
            raise RuntimeError("downloaded object is not a valid XLSX/ZIP container")

        with zipfile.ZipFile(tmp, "r") as zf:
            bad = zf.testzip()
            if bad is not None:
                raise RuntimeError(f"downloaded XLSX contains corrupt member: {bad}")

        tmp.replace(BID_FIG3)
        return BID_FIG3

    except Exception as e:
        try:
            if tmp.exists():
                tmp.unlink()
        except Exception:
            pass
        note = META / "phase2a_BID_Fig3_download_required.txt"
        note.write_text(
            "Phase 2A requires the official BID-seq Source Data Fig. 3 XLSX.\n\n"
            f"Official URL:\n{BID_FIG3_URL}\n\n"
            "Download it without renaming to:\n"
            f"{BID_FIG3}\n\n"
            f"Automatic-download error:\n{repr(e)}\n",
            encoding="utf-8",
        )
        raise RuntimeError(
            "Official BID Source Data Fig.3 is required and automatic download failed. "
            rf"Download Source Data Fig. 3 to {BID_FIG3}. "
            rf"See {note}"
        ) from e


def _gene_token_match(value: Any, gene_lookup: Dict[str, str]) -> Optional[str]:
    """
    Conservative source-data gene-name matching.
    Exact gene symbol is preferred. A tokenized match is allowed only when
    exactly one source gene symbol is present in the cell.
    """
    raw = ss(value)
    if not raw:
        return None
    up = raw.upper().strip()

    if up in gene_lookup:
        return gene_lookup[up]

    tokens = set(re.findall(r"[A-Z0-9][A-Z0-9.-]*", up))
    hits = [gene_lookup[t] for t in tokens if t in gene_lookup]
    hits = sorted(set(hits))
    return hits[0] if len(hits) == 1 else None


def _segment_gene_hits(hits: List[Tuple[int, str]], max_gap: int = 2):
    """Split gene-hit rows into panel-like contiguous blocks."""
    if not hits:
        return []
    hits = sorted(hits)
    blocks = []
    cur = [hits[0]]
    for item in hits[1:]:
        if item[0] - cur[-1][0] <= max_gap:
            cur.append(item)
        else:
            blocks.append(cur)
            cur = [item]
    blocks.append(cur)
    return blocks


def _writer_hits_in_context(matrix: List[List[Any]], r0: int, r1: int,
                            gene_col: int) -> List[str]:
    """
    Search headers/nearby cells for the eight PUS names. Fig.3k is explicitly
    a gene x PUS-knockdown heatmap, so multiple writer headers are expected.
    """
    if not matrix:
        return []
    maxc = max(len(r) for r in matrix)
    rr0 = max(0, r0 - 12)
    rr1 = min(len(matrix), r0 + 8)
    cc0 = max(0, gene_col - 2)
    cc1 = min(maxc, gene_col + 15)

    context = []
    for r in range(rr0, rr1):
        row = matrix[r]
        for c in range(cc0, cc1):
            if c < len(row):
                context.append(ss(row[c]).upper())
    text = " ".join(context)

    # Longest names first avoids PUS7 matching PUS7L as an independent hit.
    ordered = sorted(WRITERS, key=len, reverse=True)
    found = []
    for w in ordered:
        if re.search(rf"(?<![A-Z0-9]){re.escape(w)}(?![A-Z0-9])", text):
            found.append(w)
    return sorted(found)


def discover_fig3k_genes(fig3_path: Path, source_table: pd.DataFrame):
    """
    Locate the author-defined Fig.3k heatmap block.

    Crucially, membership in this block IS the published outcome:
    104 shControl HeLa mRNA Ψ sites showing reduced modification under
    one or more of eight PUS depletions.

    We do not derive membership from numeric KD values.
    """
    source_genes = sorted(set(source_table["gene"].dropna().astype(str)))
    gene_lookup = {g.upper(): g for g in source_genes}

    wb = openpyxl.load_workbook(fig3_path, read_only=True, data_only=True)
    candidates = []

    for ws in wb.worksheets:
        matrix = [list(row) for row in ws.iter_rows(values_only=True)]
        if not matrix:
            continue
        maxc = max(len(r) for r in matrix)

        for c in range(maxc):
            hits = []
            for r, row in enumerate(matrix):
                v = row[c] if c < len(row) else None
                g = _gene_token_match(v, gene_lookup)
                if g is not None:
                    hits.append((r, g))

            for block in _segment_gene_hits(hits, max_gap=2):
                if len(block) < 70:
                    continue

                r0, r1 = block[0][0], block[-1][0]
                genes = [g for _, g in block]
                writer_hits = _writer_hits_in_context(matrix, r0, r1, c)

                # Fig.3k is explicitly a 104-site x eight-PUS matrix.
                # Exact 104 rows is the primary provenance anchor.
                candidate = {
                    "sheet": ws.title,
                    "gene_col_1based": c + 1,
                    "start_row_1based": r0 + 1,
                    "end_row_1based": r1 + 1,
                    "n_gene_rows": len(block),
                    "n_unique_gene_symbols": len(set(genes)),
                    "writer_hits": ",".join(writer_hits),
                    "n_writer_hits": len(writer_hits),
                    "distance_from_104": abs(len(block) - PUBLISHED_WRITER_UNION),
                    "_genes": genes,
                    "_rows": [r for r, _ in block],
                }
                candidates.append(candidate)

    wb.close()

    audit = pd.DataFrame([
        {k: v for k, v in c.items() if not k.startswith("_")}
        for c in candidates
    ])
    audit.to_csv(
        META / "phase2a_fig3k_candidate_blocks.tsv",
        sep="\t", index=False
    )

    exact = [
        c for c in candidates
        if c["n_gene_rows"] == PUBLISHED_WRITER_UNION
        and c["n_writer_hits"] >= 4
    ]

    if len(exact) != 1:
        raise RuntimeError(
            "Could not uniquely identify the author-defined 104-site Fig.3k "
            f"heatmap block (eligible blocks={len(exact)}). "
            r"See 00_meta\phase2a_fig3k_candidate_blocks.tsv"
        )

    chosen = exact[0]
    genes = chosen["_genes"]

    # Preserve multiplicity: Fig.3k has one row per responsive site, although
    # its visible row label is the corresponding gene name.
    fig_counts = pd.Series(genes).value_counts().to_dict()
    src_counts = source_table["gene"].value_counts().to_dict()

    mapping_audit = []
    ambiguous_genes = set()
    for gene, src_n in src_counts.items():
        fig_n = int(fig_counts.get(gene, 0))
        if src_n == 1:
            status = "resolved_assigned" if fig_n == 1 else (
                "resolved_not_assigned" if fig_n == 0 else "invalid_multiplicity"
            )
        else:
            # Duplicate-gene loci can be assigned only if multiplicity itself
            # makes the answer unambiguous. Otherwise exclude the gene from the
            # source-positive analysis instead of guessing which locus is responsive.
            if fig_n == 0:
                status = "resolved_all_not_assigned"
            elif fig_n == src_n:
                status = "resolved_all_assigned"
            else:
                status = "ambiguous_duplicate_gene"
                ambiguous_genes.add(gene)

        mapping_audit.append({
            "gene": gene,
            "n_source_loci": int(src_n),
            "n_Fig3k_rows": int(fig_n),
            "mapping_status": status,
        })

    mapdf = pd.DataFrame(mapping_audit)
    mapdf.to_csv(
        META / "phase2a_fig3k_mapping_audit.tsv",
        sep="\t", index=False
    )

    # Any unique-gene multiplicity >1 in Fig3k means the block isn't a valid
    # one-row-per-site representation for our source table.
    invalid = mapdf["mapping_status"].eq("invalid_multiplicity")
    if invalid.any():
        raise RuntimeError(
            "Fig.3k candidate contains impossible multiplicity for one or more "
            "unique source genes. No outcome labels were created. "
            r"See 00_meta\phase2a_fig3k_mapping_audit.tsv"
        )

    report = {
        "source_data_file": str(fig3_path),
        "source_data_url": BID_FIG3_URL,
        "selected_sheet": chosen["sheet"],
        "gene_column_1based": chosen["gene_col_1based"],
        "start_row_1based": chosen["start_row_1based"],
        "end_row_1based": chosen["end_row_1based"],
        "published_Fig3k_rows_recovered": len(genes),
        "writer_headers_detected": chosen["writer_hits"].split(",")
            if chosen["writer_hits"] else [],
        "ambiguous_duplicate_genes_excluded": sorted(ambiguous_genes),
        "n_ambiguous_duplicate_genes": len(ambiguous_genes),
        "outcome_definition": (
            "membership in author Source Data Fig.3k 104-site heatmap; "
            "no numeric PUS-response threshold inferred"
        ),
    }
    return fig_counts, ambiguous_genes, report


def read_bid_table8() -> Tuple[pd.DataFrame, Dict]:
    """
    Table 8 supplies the 133-site source-positive universe.
    Source Data Fig.3k supplies the author-defined biological outcome.
    """
    sheet = find_table8_sheet(BID_SUPP)
    raw = pd.read_excel(BID_SUPP, sheet_name=sheet, header=3)
    raw.columns = [clean_col(c) for c in raw.columns]

    # Some copies may place the header one row earlier/later. Fall back to the
    # same matrix-based header detector used previously if needed.
    if not {"chr", "pos", "strand"}.issubset(raw.columns):
        wb0 = openpyxl.load_workbook(BID_SUPP, read_only=True, data_only=True)
        ws0 = wb0[sheet]
        matrix = [list(r) for r in ws0.iter_rows(values_only=True)]
        wb0.close()
        h = detect_header(matrix)
        header = [clean_col(v) for v in matrix[h]]
        rows = []
        for row in matrix[h+1:]:
            vals = list(row) + [None] * max(0, len(header)-len(row))
            vals = vals[:len(header)]
            if any(ss(v) for v in vals):
                rows.append(dict(zip(header, vals)))
        raw = pd.DataFrame(rows)
        header_row_1based = h + 1
    else:
        header_row_1based = 4

    required = {"chr", "pos", "strand"}
    if not required.issubset(raw.columns):
        raise RuntimeError(
            f"BID Table8 schema unresolved: columns={list(raw.columns)}"
        )

    gene_col = next(
        (c for c in ["name", "gene", "gene_name"] if c in raw.columns), None
    )
    if gene_col is None:
        raise RuntimeError("BID Table8 gene column unresolved.")

    x = pd.DataFrame({
        "chrom": raw["chr"].map(norm_chr),
        "pos1": raw["pos"].map(as_int),
        "strand": raw["strand"].map(norm_strand),
        "gene": raw[gene_col].map(ss),
    })
    x = x[
        x["chrom"].ne("") & x["pos1"].notna()
        & x["strand"].ne("") & x["gene"].ne("")
    ].copy()
    x["base_locus_id"] = (
        x["chrom"] + ":" + x["strand"] + ":"
        + x["pos1"].astype("Int64").astype(str)
    )

    if len(x) != 133 or x["base_locus_id"].nunique() != 133:
        raise RuntimeError(
            f"BID Table8 gate failed: rows={len(x)}, "
            f"unique loci={x['base_locus_id'].nunique()}, expected 133/133."
        )

    fig3_path = ensure_bid_fig3_source()
    fig_counts, ambiguous_genes, fig3_report = discover_fig3k_genes(
        fig3_path, x
    )

    src_counts = x["gene"].value_counts().to_dict()
    assigned = []
    resolved = []
    for gene in x["gene"]:
        if gene in ambiguous_genes:
            assigned.append(np.nan)
            resolved.append(False)
            continue
        src_n = int(src_counts[gene])
        fig_n = int(fig_counts.get(gene, 0))
        if fig_n == 0:
            assigned.append(0)
        elif fig_n == src_n:
            assigned.append(1)
        elif src_n == 1 and fig_n == 1:
            assigned.append(1)
        else:
            assigned.append(np.nan)
            resolved.append(False)
            continue
        resolved.append(True)

    x["writer_assigned"] = assigned
    x["assignment_resolved"] = resolved
    x["writer_set"] = np.where(
        x["writer_assigned"].eq(1),
        "published_any_PUS_reduced_Fig3k",
        ""
    )

    # We deliberately do not force 104 locus labels when duplicate gene names
    # cannot identify which of two source loci corresponds to a heatmap row.
    # Instead those loci are transparently excluded from the primary regression.
    resolved_x = x[x["assignment_resolved"]].copy()

    report = {
        "sheet": sheet,
        "header_row_1based": header_row_1based,
        "unique_source_loci": int(x["base_locus_id"].nunique()),
        "source_unique_gene_symbols": int(x["gene"].nunique()),
        "published_union_expected": 104,
        "published_union_recovered_in_Fig3k": int(sum(fig_counts.values())),
        "resolved_source_loci_for_analysis": int(len(resolved_x)),
        "resolved_writer_assigned_loci": int(resolved_x["writer_assigned"].sum()),
        "unresolved_duplicate_gene_loci_excluded": int((~x["assignment_resolved"]).sum()),
        **fig3_report,
    }

    x.to_csv(
        META / "phase2a_BID_source_outcome_mapping.tsv",
        sep="\t", index=False
    )
    return x, report


# -----------------------------------------------------------------------------
# chemistry-aware local U-run intervals
# -----------------------------------------------------------------------------

def run_interval(chrom: str, strand: str, pos1: int, motif: str) -> Dict[str, Any]:
    m = motif5(motif)
    if not m:
        return {"run_start": np.nan, "run_end": np.nan, "run_complete": False,
                "single_u": False, "run_len_local": np.nan}
    l = r = 2
    while l > 0 and m[l-1] == "U":
        l -= 1
    while r < 4 and m[r+1] == "U":
        r += 1

    offsets = list(range(l-2, r-2+1))
    coords = [pos1 + o if strand == "+" else pos1 - o for o in offsets]
    complete = (l > 0 and r < 4)
    return {
        "run_start": min(coords),
        "run_end": max(coords),
        "run_complete": bool(complete),
        "single_u": bool(l == 2 and r == 2),
        "run_len_local": int(r-l+1),
    }


def attach_run_fields(df: pd.DataFrame) -> pd.DataFrame:
    x = df.copy()
    recs = [
        run_interval(r.chrom, r.strand, int(r.pos1), r.motif)
        for r in x.itertuples()
    ]
    rr = pd.DataFrame(recs, index=x.index)
    return pd.concat([x, rr], axis=1)


def external_bacs() -> pd.DataFrame:
    x = pd.read_csv(BACS, sep="\t", low_memory=False)
    y = pd.DataFrame({
        "chrom": x["chrom"].map(norm_chr),
        "pos1": pd.to_numeric(x["pos1"], errors="coerce").astype("Int64"),
        "strand": x["strand"].map(norm_strand),
        "motif": x["motif_5mer_reported"].map(motif5),
    }).dropna(subset=["pos1"])
    y = y[y["motif"].ne("") & y["strand"].ne("")].copy()
    y["pos1"] = y["pos1"].astype(int)
    return attach_run_fields(y)


def find_elap_seqcols(x: pd.DataFrame) -> List[str]:
    raw = [c for c in x.columns if c.startswith("raw__")]
    cands = []
    for i in range(len(raw)-4):
        cols = raw[i:i+5]
        frac = []
        for c in cols:
            z = x[c].fillna("").astype(str).str.upper().str.strip()
            frac.append(z.isin(["A","C","G","T","U"]).mean())
        center = x[cols[2]].fillna("").astype(str).str.upper().str.strip()
        center_u = center.isin(["T","U"]).mean()
        if min(frac) > .90 and center_u > .90:
            cands.append((np.mean(frac)+center_u, cols))
    if not cands:
        raise RuntimeError("ELAP five-base sequence window could not be resolved.")
    return sorted(cands, reverse=True)[0][1]


def external_elap() -> pd.DataFrame:
    x = pd.read_csv(ELAP, sep="\t", low_memory=False)
    cols = find_elap_seqcols(x)
    seq = x[cols].fillna("").astype(str).agg("".join, axis=1)
    y = pd.DataFrame({
        "chrom": x["chrom"].map(norm_chr),
        "pos1": pd.to_numeric(x["pos1"], errors="coerce").astype("Int64"),
        "strand": x["strand"].map(norm_strand),
        "motif": seq.map(motif5),
    }).dropna(subset=["pos1"])
    y = y[y["motif"].ne("") & y["strand"].ne("")].copy()
    y["pos1"] = y["pos1"].astype(int)
    return attach_run_fields(y)


def build_support_index(ext: pd.DataFrame):
    exact = set(zip(ext["chrom"], ext["strand"], ext["pos1"]))
    complete_runs = set(
        zip(
            ext.loc[ext["run_complete"], "chrom"],
            ext.loc[ext["run_complete"], "strand"],
            ext.loc[ext["run_complete"], "run_start"].astype(int),
            ext.loc[ext["run_complete"], "run_end"].astype(int),
        )
    )
    return exact, complete_runs


def support_for_source(src: pd.DataFrame, ext: pd.DataFrame, prefix: str) -> pd.DataFrame:
    exact, runs = build_support_index(ext)
    out = src.copy()
    exact_hits = []
    equiv_hits = []
    modes = []

    for r in out.itertuples():
        ex = (r.chrom, r.strand, int(r.pos1)) in exact
        eq = False
        if bool(r.run_complete):
            eq = (
                r.chrom, r.strand, int(r.run_start), int(r.run_end)
            ) in runs
        hit = ex or eq
        exact_hits.append(int(ex))
        equiv_hits.append(int(hit))
        if ex:
            modes.append("exact")
        elif eq:
            modes.append("complete_U_run")
        else:
            modes.append("none")

    out[f"{prefix}_support_exact"] = exact_hits
    out[f"{prefix}_support"] = equiv_hits
    out[f"{prefix}_match_mode"] = modes
    return out


# -----------------------------------------------------------------------------
# motif flags / models
# -----------------------------------------------------------------------------

def known_motif_flag(m: str) -> int:
    return int(bool(PUS7_RE.match(m)) or bool(TRUB1_RE.match(m)))


def logistic_cluster(df: pd.DataFrame, outcome: str, predictors: List[str],
                     cluster: str = "motif") -> pd.DataFrame:
    d = df[[outcome, cluster] + predictors].copy()
    for c in [outcome] + predictors:
        d[c] = pd.to_numeric(d[c], errors="coerce")
    d = d.dropna()

    y = d[outcome].to_numpy(float)
    X = np.column_stack([np.ones(len(d))] + [d[p].to_numpy(float) for p in predictors])
    names = ["intercept"] + predictors

    beta = np.zeros(X.shape[1])
    for _ in range(200):
        eta = np.clip(X @ beta, -30, 30)
        pr = 1/(1+np.exp(-eta))
        W = np.clip(pr*(1-pr), 1e-8, None)
        H = X.T @ (W[:,None]*X)
        score = X.T @ (y-pr)
        step = np.linalg.pinv(H) @ score
        beta2 = beta + step
        if np.max(np.abs(step)) < 1e-9:
            beta = beta2
            break
        beta = beta2

    eta = np.clip(X @ beta, -30, 30)
    pr = 1/(1+np.exp(-eta))
    W = np.clip(pr*(1-pr), 1e-8, None)
    bread = np.linalg.pinv(X.T @ (W[:,None]*X))

    groups = defaultdict(list)
    for i,g in enumerate(d[cluster].astype(str)):
        groups[g].append(i)
    meat = np.zeros((X.shape[1],X.shape[1]))
    for ids in groups.values():
        sg = X[ids].T @ (y[ids]-pr[ids])
        meat += np.outer(sg,sg)

    G = len(groups)
    N,K = X.shape
    correction = (G/max(G-1,1))*((N-1)/max(N-K,1))
    cov = bread @ meat @ bread * correction
    se = np.sqrt(np.clip(np.diag(cov),0,None))

    rows = []
    for name,b,s0 in zip(names,beta,se):
        z = b/s0 if s0 > 0 else np.nan
        p = math.erfc(abs(z)/math.sqrt(2)) if np.isfinite(z) else np.nan
        rows.append({
            "term": name, "beta": b, "SE_motif_cluster": s0, "z": z, "p": p,
            "OR": math.exp(b),
            "OR_CI_low": math.exp(b-1.96*s0),
            "OR_CI_high": math.exp(b+1.96*s0),
            "n": N, "n_motif_clusters": G,
        })
    return pd.DataFrame(rows)


def gradient_table(df: pd.DataFrame, outcome: str, breadth: str) -> pd.DataFrame:
    rows = []
    for c in sorted(df[breadth].dropna().unique()):
        z = df[df[breadth] == c]
        n = len(z)
        k = int(z[outcome].sum())
        lo,hi = wilson(k,n)
        rows.append({
            "chemistry_breadth": int(c),
            "n": n,
            "supported_or_assigned": k,
            "rate": k/n if n else np.nan,
            "wilson95_low": lo,
            "wilson95_high": hi,
        })
    return pd.DataFrame(rows)


# -----------------------------------------------------------------------------
# DRS held-out parser
# -----------------------------------------------------------------------------

LOC_RE = re.compile(r"(chr(?:[0-9]+|X|Y|M|MT))[:_](\d+)", re.I)


def truthy(v: Any) -> bool:
    z = ss(v).lower()
    if z in {"", "0", "0.0", "false", "no", "na", "n/a", "nan"}:
        return False
    return True


def parse_drs_hela() -> Tuple[Set[Tuple[str,int]], Dict]:
    if not DRS.exists():
        return set(), {"status": "DRS_NORMALIZED_TABLE_MISSING", "path": str(DRS)}

    df = pd.read_csv(DRS, sep="\t", dtype=str, keep_default_na=False, low_memory=False)

    # Find HeLa rows in row-wise representation.
    row_hela = pd.Series(False, index=df.index)
    for c in df.columns:
        if c == "source_file":
            continue
        vals = df[c].astype(str)
        row_hela |= vals.str.contains(r"(?i)\bHeLa\b", regex=True)

    # If a HeLa-named column exists, require that column to be truthy.
    hela_cols = [c for c in df.columns if re.search(r"(?i)\bhela\b", c)]
    for c in hela_cols:
        row_hela |= df[c].map(truthy)

    cand = df[row_hela].copy()

    # First preference: a combined genomic-locus field.
    loci = set()
    locus_columns = []
    for c in cand.columns:
        vals = cand[c].astype(str)
        matches = vals.map(lambda z: LOC_RE.search(z))
        n = int(matches.notna().sum())
        if n >= max(20, int(.10*max(len(cand),1))):
            locus_columns.append((n,c))
    if locus_columns:
        locus_columns.sort(reverse=True)
        c = locus_columns[0][1]
        for v in cand[c]:
            m = LOC_RE.search(str(v))
            if m:
                loci.add((norm_chr(m.group(1)), int(m.group(2))))
        return loci, {
            "status": "RESOLVED_COMBINED_LOCUS",
            "n_total_rows": len(df),
            "n_HeLa_candidate_rows": len(cand),
            "coordinate_column": c,
            "n_HeLa_loci": len(loci),
            "strand_used": False,
        }

    # Second preference: separate chromosome and genomic position columns.
    chrom_scores = []
    pos_scores = []
    for c in cand.columns:
        vals = cand[c].astype(str)
        chrfrac = vals.map(lambda z: bool(re.fullmatch(r"chr(?:\d+|X|Y|M|MT)", z, re.I))).mean()
        if chrfrac > .5:
            chrom_scores.append((chrfrac, c))
        numfrac = pd.to_numeric(vals, errors="coerce").notna().mean()
        name_bonus = 1 if re.search(r"(?i)(genomic|position|pos|coordinate)", c) else 0
        if numfrac > .5 and name_bonus:
            pos_scores.append((name_bonus, numfrac, c))

    if chrom_scores and pos_scores:
        chrom_col = sorted(chrom_scores, reverse=True)[0][1]
        pos_col = sorted(pos_scores, reverse=True)[0][2]
        for _,r in cand.iterrows():
            chrom = norm_chr(r[chrom_col])
            pos = as_int(r[pos_col])
            if chrom and pos:
                loci.add((chrom,pos))
        if len(loci) >= 20:
            return loci, {
                "status": "RESOLVED_SEPARATE_COLUMNS",
                "n_total_rows": len(df),
                "n_HeLa_candidate_rows": len(cand),
                "chrom_column": chrom_col,
                "position_column": pos_col,
                "n_HeLa_loci": len(loci),
                "strand_used": False,
            }

    # Preserve useful diagnostic without guessing.
    preview = cand.head(100) if len(cand) else df.head(100)
    preview.to_csv(META / "phase2a_DRS_S6_parser_preview.tsv", sep="\t", index=False)
    return set(), {
        "status": "UNRESOLVED_NO_GUESS",
        "n_total_rows": len(df),
        "n_HeLa_candidate_rows": len(cand),
        "columns": list(df.columns),
        "preview": r"00_meta\phase2a_DRS_S6_parser_preview.tsv",
    }


def drs_support_for_source(src: pd.DataFrame, drs_loci: Set[Tuple[str,int]]) -> pd.Series:
    out = []
    by_chr = defaultdict(set)
    for chrom,pos in drs_loci:
        by_chr[chrom].add(pos)

    for r in src.itertuples():
        positions = by_chr.get(r.chrom, set())
        if int(r.pos1) in positions:
            out.append(1)
        elif bool(r.run_complete) and any(
            int(r.run_start) <= p <= int(r.run_end) for p in positions
        ):
            out.append(1)
        else:
            out.append(0)
    return pd.Series(out, index=src.index)


# -----------------------------------------------------------------------------
# BACS KO discovery for next independent replication
# -----------------------------------------------------------------------------

def audit_bacs_ko_candidates() -> pd.DataFrame:
    if not BACS_BOOK.exists():
        return pd.DataFrame([{"status":"BACS_WORKBOOK_MISSING"}])

    wb = openpyxl.load_workbook(BACS_BOOK, read_only=True, data_only=True)
    rows = []
    for ws in wb.worksheets:
        sample = []
        for i,row in enumerate(ws.iter_rows(values_only=True), start=1):
            sample.append(list(row))
            if i >= 20:
                break
        text = " ".join(ss(v) for row in sample for v in row)
        hits = [w for w in ["TRUB1","PUS7","PUS1","KO","knockout","dependent"] if w.lower() in text.lower()]
        if hits:
            rows.append({
                "sheet": ws.title,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "hits": ",".join(hits),
                "preview": " | ".join(
                    ss(v) for row in sample[:5] for v in row if ss(v)
                )[:1500],
            })
    wb.close()
    return pd.DataFrame(rows)


# -----------------------------------------------------------------------------
# main
# -----------------------------------------------------------------------------

def main():
    print("="*100)
    print("TRACE PHASE 2A — CHEMISTRY-DIVERSE SUPPORT -> BIOLOGICAL EVIDENCE")
    print("="*100)

    for p in [BID_SH, BID_SUPP, BACS, ELAP]:
        if not p.exists():
            raise RuntimeError(f"Missing required input: {p}")

    source = read_bid_shcontrol()
    table8, table8_report = read_bid_table8()

    source = source.merge(
        table8[["base_locus_id","writer_assigned","writer_set","assignment_resolved"]],
        on="base_locus_id", how="left", validate="one_to_one"
    )
    if source["assignment_resolved"].isna().any():
        n = int(source["assignment_resolved"].isna().sum())
        raise RuntimeError(
            f"{n} BID shControl loci failed the baseline Table8 join."
        )

    # Conservative primary universe: exclude only duplicate-gene loci for which
    # Fig.3k's gene-only row label cannot identify the responsive nucleotide.
    n_before_outcome_resolution = len(source)
    source = source[source["assignment_resolved"]].copy()
    source["writer_assigned"] = source["writer_assigned"].astype(int)

    if len(source) < 120:
        raise RuntimeError(
            f"Only {len(source)}/133 source loci have unambiguous Fig3k outcome "
            "mapping; this is too much attrition for the predeclared analysis."
        )

    source["known_writer_motif"] = source["motif"].map(known_motif_flag)
    source = attach_run_fields(source)

    bacs = external_bacs()
    elap = external_elap()

    source = support_for_source(source, bacs, "BACS")
    source = support_for_source(source, elap, "ELAP")

    source["chemistry_breadth"] = source["BACS_support"] + source["ELAP_support"]
    source["chemistry_breadth_exact"] = (
        source["BACS_support_exact"] + source["ELAP_support_exact"]
    )

    # Primary biological-validation tables.
    pus_gradient = gradient_table(source, "writer_assigned", "chemistry_breadth")
    pus_gradient.to_csv(AIM2 / "aim2a_PUS_gradient.tsv", sep="\t", index=False)

    model_primary = logistic_cluster(
        source, "writer_assigned",
        ["chemistry_breadth", "source_fraction_rank"]
    )
    model_primary["model"] = "primary_run_aware_adjust_source_fraction"

    model_motif = logistic_cluster(
        source, "writer_assigned",
        ["chemistry_breadth", "source_fraction_rank", "known_writer_motif"]
    )
    model_motif["model"] = "sensitivity_add_known_writer_motif"

    model_exact = logistic_cluster(
        source, "writer_assigned",
        ["chemistry_breadth_exact", "source_fraction_rank"]
    )
    model_exact["model"] = "sensitivity_exact_coordinate_only"

    single = source[source["single_u"]].copy()
    model_single = logistic_cluster(
        single, "writer_assigned",
        ["chemistry_breadth", "source_fraction_rank"]
    )
    model_single["model"] = "sensitivity_single_U_only"

    models = pd.concat(
        [model_primary, model_motif, model_exact, model_single],
        ignore_index=True
    )
    models.to_csv(AIM2 / "aim2a_PUS_adjusted_models.tsv", sep="\t", index=False)

    # Held-out DRS.
    drs_loci, drs_report = parse_drs_hela()
    if drs_loci:
        source["DRS_reported_support"] = drs_support_for_source(source, drs_loci)
        drs_gradient = gradient_table(
            source, "DRS_reported_support", "chemistry_breadth"
        )
        drs_model = logistic_cluster(
            source, "DRS_reported_support",
            ["chemistry_breadth", "source_fraction_rank"]
        )
        drs_model["model"] = "heldout_DRS_run_aware"
    else:
        source["DRS_reported_support"] = np.nan
        drs_gradient = pd.DataFrame()
        drs_model = pd.DataFrame()

    drs_gradient.to_csv(AIM2 / "aim2a_heldout_DRS_gradient.tsv", sep="\t", index=False)
    drs_model.to_csv(AIM2 / "aim2a_heldout_DRS_model.tsv", sep="\t", index=False)

    source.to_csv(AIM2 / "aim2a_source_positive_evidence_table.tsv", sep="\t", index=False)

    # Prepare the BACS KO replication efficiently without analyzing guessed schema.
    bacs_ko = audit_bacs_ko_candidates()
    bacs_ko.to_csv(META / "phase2a_BACS_KO_candidate_sheets.tsv", sep="\t", index=False)

    # Compact interpretation components. No arbitrary combined score.
    primary_beta = model_primary.loc[
        model_primary["term"] == "chemistry_breadth"
    ].iloc[0].to_dict()

    exact_beta = model_exact.loc[
        model_exact["term"] == "chemistry_breadth_exact"
    ].iloc[0].to_dict()

    single_beta = model_single.loc[
        model_single["term"] == "chemistry_breadth"
    ].iloc[0].to_dict() if len(model_single) else {}

    drs_beta = {}
    if len(drs_model):
        drs_beta = drs_model.loc[
            drs_model["term"] == "chemistry_breadth"
        ].iloc[0].to_dict()

    status = (
        "AIM2A_PLUS_HELDOUT_DRS_RESULT_READY"
        if drs_loci else
        "AIM2A_PUS_RESULT_READY_DRS_SCHEMA_REVIEW"
    )

    contract = {
        "phase": "2A",
        "scientific_story": "Calibration -> Measurement -> Evidence -> Inference",
        "question": (
            "Within a fixed source-positive BID shControl HeLa universe, does "
            "support from independent chemistry families carry additional "
            "biological evidence?"
        ),
        "source_universe": (
            "133 BID-seq HeLa shControl mRNA Ψ sites (>10% published fraction) "
            "used by the BID paper for eight-PUS perturbation analysis."
        ),
        "biological_outcome": (
            "Membership in the author-defined Source Data Fig.3k 104-site heatmap "
            "of sites with reduced modification under >=1 tested PUS depletion. "
            "Duplicate-gene loci are excluded if the gene-only Fig3k row label "
            "cannot identify the nucleotide. Complement means not included in "
            "this 104-site published response set, not biological negative."
        ),
        "external_evidence": {
            "BACS": "HeLa cyclization/conversion",
            "ELAP": "HeLa enzymatic labeling/RT-stop",
            "DRS": "HeLa native-RNA, held out until after evidence strata are frozen",
        },
        "chemistry_breadth": "BACS_support + ELAP_support; range 0-2",
        "matching": (
            "exact base, or identical fully-resolved local contiguous-U interval; "
            "edge-truncated local runs fall back to exact matching; no +/-k."
        ),
        "primary_model": (
            "writer_assigned ~ chemistry_breadth + BID source fraction percentile; "
            "motif-cluster robust SE"
        ),
        "sensitivities": [
            "add known TRUB1/PUS7 motif indicator",
            "exact-coordinate-only matching",
            "single-U source sites only",
        ],
        "heldout_DRS_interpretation": (
            "DRS reported support gradient only. Absence from DRS call table is "
            "never interpreted as unmodified or false."
        ),
        "no_truth_score": True,
    }
    (META / "phase2a_analysis_contract.json").write_text(
        json.dumps(contract, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    summary = {
        "status": status,
        "BID_table8_provenance_gate": table8_report,
        "n_source_positive_published": 133,
        "n_source_positive_analyzable": int(len(source)),
        "n_outcome_ambiguous_loci_excluded": int(n_before_outcome_resolution - len(source)),
        "n_PUS_assigned_analyzable": int(source["writer_assigned"].sum()),
        "source_single_U": int(source["single_u"].sum()),
        "evidence_counts": source["chemistry_breadth"].value_counts().sort_index().to_dict(),
        "BACS_support_n": int(source["BACS_support"].sum()),
        "ELAP_support_n": int(source["ELAP_support"].sum()),
        "PUS_response_gradient": pus_gradient.to_dict(orient="records"),
        "primary_chemistry_breadth_effect": primary_beta,
        "exact_only_effect": exact_beta,
        "single_U_effect": single_beta,
        "DRS_parser": drs_report,
        "DRS_support_n": (
            int(source["DRS_reported_support"].sum())
            if source["DRS_reported_support"].notna().any() else None
        ),
        "DRS_gradient": drs_gradient.to_dict(orient="records"),
        "DRS_chemistry_breadth_effect": drs_beta,
        "BACS_KO_candidate_sheets_n": int(len(bacs_ko)),
        "next_gate": (
            "Interpret PUS and held-out DRS evidence together. If chemistry breadth "
            "shows a consistent positive association, replicate the biological "
            "validation in the independent BACS KO source-positive dataset using "
            "the candidate sheets already audited here; then freeze Aim2. If the "
            "association is null, do not change thresholds—report the null and use "
            "DRS/PUS patterns to decide whether evidence breadth is biologically "
            "informative."
        ),
    }
    (META / "phase2a_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2, default=str),
        encoding="utf-8"
    )

    print()
    print("STATUS:", status)
    print("BID published source-positive sites: 133")
    print("Analyzable after duplicate-gene outcome resolution:", len(source))
    print("PUS-assigned in analyzable subset:", int(source["writer_assigned"].sum()))
    print("Chemistry breadth counts:", source["chemistry_breadth"].value_counts().sort_index().to_dict())
    print()
    print("PUS RESPONSE GRADIENT")
    print(pus_gradient.to_string(index=False))
    print()
    print("PRIMARY ADJUSTED MODEL")
    print(model_primary.to_string(index=False))
    print()
    print("DRS:", drs_report)
    if len(drs_gradient):
        print(drs_gradient.to_string(index=False))
        print(drs_model.to_string(index=False))
    print()
    print(r"Return/upload:")
    print(r"  D:\RNA\Trace\00_meta\phase2a_summary.json")
    print(r"  D:\RNA\Trace\00_meta\phase2a_analysis_contract.json")
    print(r"  D:\RNA\Trace\00_meta\phase2a_BACS_KO_candidate_sheets.tsv")
    print(r"  D:\RNA\Trace\05_aim2\aim2a_PUS_gradient.tsv")
    print(r"  D:\RNA\Trace\05_aim2\aim2a_PUS_adjusted_models.tsv")
    print(r"  D:\RNA\Trace\05_aim2\aim2a_heldout_DRS_gradient.tsv")
    print(r"  D:\RNA\Trace\05_aim2\aim2a_heldout_DRS_model.tsv")
    print(r"  D:\RNA\Trace\05_aim2\aim2a_source_positive_evidence_table.tsv")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception:
        err = traceback.format_exc()
        (LOGS / "phase2a_FATAL_ERROR.txt").write_text(err, encoding="utf-8")
        print("\nFATAL ERROR\n")
        print(err)
        print(r"Send D:\RNA\Trace\logs\phase2a_FATAL_ERROR.txt back to ChatGPT.")
        try:
            input("\nPress Enter to close...")
        except Exception:
            pass
        raise
